"""Parser for the locally verified eight-column Khan Bank MNT export.

Numeric XML lexemes feed Decimal directly so openpyxl float conversion cannot
silently erase precision. openpyxl supplies workbook structure and cached values.
No statement contents, paths, or exception details are logged.
"""

import posixpath
import re
import warnings
from contextlib import ExitStack
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.models import ImportRow, money_to_minor, require_integer
from app.parsers.base import (
    ParseIssue, ParseResult, ParsedTransaction, ReconciliationCheck, StatementMetadata,
)


HEADERS = (
    "Гүйлгээний огноо", "Салбар", "Эхний үлдэгдэл", "Кредит гүйлгээ",
    "Дебит гүйлгээ", "Эцсийн үлдэгдэл", "Гүйлгээний утга", "Харьцсан данс",
)
_TITLE_SUFFIX = " дансны дэлгэрэнгүй хуулга"
_TOTAL_LABEL = "Нийт дүн:"
_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_MONEY_TEXT = re.compile(r"-?(?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(?:\.[0-9]{1,2})?")
_DATE_TIME = r"[0-9]{4}-[0-9]{2}-[0-9]{2} [0-9]{2}:[0-9]{2}:[0-9]{2}"


class _InvalidCell(Exception):
    def __init__(self, code: str, message: str, column: str | None = None) -> None:
        self.code, self.message, self.column = code, message, column


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _label(value: object) -> str:
    return " ".join(value.split()) if isinstance(value, str) else ""


def _header(sheet, original, row: int) -> bool:
    return tuple(_label(sheet.cell(row, c).value) for c in range(1, 9)) == HEADERS and all(
        _blank(sheet.cell(row, c).value) and _blank(original.cell(row, c).value)
        for c in range(9, max(sheet.max_column, original.max_column) + 1)
    )


def _numeric_lexemes(payload: bytes, sheet_title: str) -> dict[str, str]:
    """Read selected sheet numbers/cached numbers without binary arithmetic."""
    with ZipFile(BytesIO(payload)) as package:
        workbook = ET.fromstring(package.read("xl/workbook.xml"))
        sheets = workbook.findall(f"{{{_NS}}}sheets/{{{_NS}}}sheet")
        sheet = next(s for s in sheets if s.get("name") == sheet_title)
        relationship_id = sheet.attrib[f"{{{_REL}}}id"]
        relationships = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        relationship = next(r for r in relationships.findall(f"{{{_PACKAGE_REL}}}Relationship")
                            if r.get("Id") == relationship_id)
        if relationship.get("TargetMode") == "External":
            raise ValueError("External worksheets are unsupported.")
        target = relationship.attrib["Target"]
        name = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        if not name.startswith("xl/"):
            raise ValueError("Unsupported worksheet relationship.")
        xml_sheet = ET.fromstring(package.read(name))
        numbers = {}
        for cell in xml_sheet.findall(f"{{{_NS}}}sheetData/{{{_NS}}}row/{{{_NS}}}c"):
            value = cell.find(f"{{{_NS}}}v")
            if cell.get("t", "n") == "n" and value is not None and value.text is not None:
                numbers[cell.attrib["r"]] = value.text
        return numbers


class _Cells:
    def __init__(self, original, cached, numbers: dict[str, str]) -> None:
        self.original, self.cached, self.numbers = original, cached, numbers

    def value(self, row: int, column: int):
        source = self.original.cell(row, column)
        cached = self.cached.cell(row, column)
        if source.data_type == "f" and cached.value is None:
            raise _InvalidCell("missing_formula_cache", "Formula has no usable cached value.", source.column_letter)
        if cached.data_type == "e":
            raise _InvalidCell("cell_error", "Cell contains a spreadsheet error.", source.column_letter)
        return cached.value

    def text(self, row: int, column: int, *, optional: bool = False) -> str | None:
        value = self.value(row, column)
        if optional and _blank(value):
            return None
        if not isinstance(value, str) or not value.strip():
            raise _InvalidCell("invalid_text", "Required text is absent or is not stored as text.",
                               self.original.cell(row, column).column_letter)
        return value

    def money(self, row: int, column: int) -> int:
        value = self.value(row, column)
        cell = self.cached.cell(row, column)
        try:
            if type(value) in (int, float):
                # The OOXML text, not this float, is used for all money arithmetic.
                number = Decimal(self.numbers[cell.coordinate])
            elif isinstance(value, str) and _MONEY_TEXT.fullmatch(value.strip()):
                number = Decimal(value.strip().replace(",", ""))
            else:
                raise ValueError("Unsupported money cell.")
            return money_to_minor(number)
        except (ValueError, InvalidOperation, KeyError, OverflowError):
            raise _InvalidCell("invalid_money", "Money must have a verified decimal format and at most two decimal places.",
                               cell.column_letter) from None


def _date(value: str, pattern: str, format_string: str, column: str) -> str:
    if re.fullmatch(pattern, value) is None:
        raise _InvalidCell("invalid_date", "Date does not match the verified format.", column)
    try:
        return datetime.strptime(value, format_string).isoformat()
    except ValueError:
        raise _InvalidCell("invalid_date", "Date contains invalid calendar or time components.", column) from None


class KhanParser:
    version = "khan-mnt-1"

    def detect(self, path: str | Path) -> bool:
        return self.parse(path).is_valid

    def parse(self, path: str | Path) -> ParseResult:
        path = Path(path)
        if path.suffix.lower() != ".xlsx":
            return ParseResult(issues=(ParseIssue("unsupported_file", "Only .xlsx workbooks are supported."),))
        try:
            payload = path.read_bytes()
            with ExitStack() as stack, warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                original = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
                stack.callback(original.close)
                cached = load_workbook(BytesIO(payload), data_only=True, keep_links=False)
                stack.callback(cached.close)
                return self._parse_workbook(original, cached, payload)
        except (OSError, BadZipFile, InvalidFileException, ET.ParseError, ValueError,
                KeyError, StopIteration, TypeError, IndexError, OverflowError):
            return ParseResult(issues=(ParseIssue("unreadable_workbook", "Workbook could not be read as the supported XLSX format."),))

    def _parse_workbook(self, original, cached, payload: bytes) -> ParseResult:
        candidates = []
        for index, sheet in enumerate(cached.worksheets):
            headers = [row for row in range(1, sheet.max_row + 1)
                       if _header(sheet, original[sheet.title], row)]
            if headers:
                candidates.append((index, sheet, headers[0]))
        if not candidates:
            return ParseResult(issues=(ParseIssue("unsupported_headers", "The verified eight-column header was not found."),))
        if len(candidates) != 1:
            return ParseResult(issues=(ParseIssue("ambiguous_worksheets", "More than one worksheet matches the verified header."),))
        index, sheet, header_row = candidates[0]
        source_sheet = original[sheet.title]
        if source_sheet.merged_cells.ranges:
            return ParseResult(issues=(ParseIssue("unsupported_merged_cells", "Merged-cell layouts have not been verified.", sheet_index=index),))
        cells = _Cells(source_sheet, sheet, _numeric_lexemes(payload, sheet.title))
        issues: list[ParseIssue] = []

        def issue(error: _InvalidCell, row: int | None = None) -> None:
            issues.append(ParseIssue(error.code, error.message, sheet_index=index,
                                     row_number=row, column=error.column))

        metadata_rows = [row for row in range(1, header_row)
                         if isinstance(sheet.cell(row, 3).value, str)
                         and sheet.cell(row, 3).value.endswith(_TITLE_SUFFIX)]
        if len(metadata_rows) != 1:
            return ParseResult(issues=(ParseIssue("invalid_metadata", "Exactly one verified title/metadata row is required before the header.", sheet_index=index),))
        metadata_row = metadata_rows[0]
        try:
            metadata = self._metadata(cells, metadata_row, sheet.title)
        except _InvalidCell as error:
            issue(error, metadata_row)
            return ParseResult(issues=tuple(issues))
        for row in range(1, header_row):
            if row != metadata_row and not self._row_blank(cells, row):
                issue(_InvalidCell("unexpected_title_row", "Unrecognized nonblank row before the transaction header."), row)

        records: list[ParsedTransaction] = []
        candidates_count = 0
        totals = None
        total_row = None
        for row in range(header_row + 1, sheet.max_row + 1):
            if self._row_blank(cells, row):
                continue
            if total_row is not None:
                issue(_InvalidCell("unexpected_after_totals", "Nonblank content appears after the final totals row."), row)
                continue
            if _header(sheet, source_sheet, row):
                issues.append(ParseIssue("repeated_header", "Exact repeated header skipped.", "warning", index, row))
                continue
            if _label(sheet.cell(row, 1).value) == _TOTAL_LABEL:
                total_row = row
                try:
                    totals = self._totals(cells, row)
                except _InvalidCell as error:
                    issue(error, row)
                continue
            candidates_count += 1
            try:
                records.append(self._transaction(cells, row, metadata))
            except _InvalidCell as error:
                issue(error, row)
        if total_row is None:
            issue(_InvalidCell("missing_totals", "The verified final totals row is missing."))
        if not records:
            issue(_InvalidCell("no_transactions", "No usable transaction rows were found."))

        checks = [ReconciliationCheck("transaction_rows", "passed" if len(records) == candidates_count else "failed")]
        check_names = ("row_balances", "running_balances", "credit_total", "debit_total", "statement_balance", "date_order")
        if any(i.severity == "error" for i in issues) or totals is None:
            checks.extend(ReconciliationCheck(name, "not_checked") for name in check_names)
        else:
            credit_total, debit_magnitude = totals
            failing_rows = {
                "row_balances": [r.source.row_number for r in records
                                 if r.opening_balance_minor + r.source.amount_minor != r.source.balance_minor],
                "running_balances": [current.source.row_number for previous, current in zip(records, records[1:])
                                     if current.opening_balance_minor != previous.source.balance_minor],
                "date_order": [current.source.row_number for previous, current in zip(records, records[1:])
                               if previous.source.occurred_at > current.source.occurred_at],
            }
            outcomes = (
                not failing_rows["row_balances"],
                not failing_rows["running_balances"],
                sum(r.credit_minor for r in records) == credit_total,
                -sum(r.debit_minor for r in records) == debit_magnitude,
                records[0].opening_balance_minor + credit_total - debit_magnitude == records[-1].source.balance_minor,
                not failing_rows["date_order"],
            )
            for name, passed in zip(check_names, outcomes):
                checks.append(ReconciliationCheck(name, "passed" if passed else "failed"))
                if not passed:
                    for row in failing_rows.get(name, [total_row]):
                        issue(_InvalidCell("reconciliation_" + name,
                                           "Statement reconciliation check failed: " + name + "."), row)
        return ParseResult(metadata=metadata, records=tuple(records), issues=tuple(issues),
                           checks=tuple(checks), candidate_row_count=candidates_count)

    @staticmethod
    def _row_blank(cells: _Cells, row: int) -> bool:
        # A formula with a missing cache is still a populated source cell.
        return all(_blank(cells.original.cell(row, c).value) and _blank(cells.cached.cell(row, c).value)
                   for c in range(1, max(cells.original.max_column, cells.cached.max_column) + 1))

    def _metadata(self, cells: _Cells, row: int, worksheet: str) -> StatementMetadata:
        for column, expected in ((1, "Хэвлэсэн огноо:"), (4, "Хэрэглэгч:"), (6, "Интервал:"), (8, "-"), (10, "Валют:")):
            if _label(cells.value(row, column)) != expected:
                raise _InvalidCell("invalid_metadata", "Metadata labels do not match the verified layout.")
        if any(not _blank(cells.original.cell(row, c).value) for c in range(13, cells.original.max_column + 1)):
            raise _InvalidCell("invalid_metadata", "Unexpected populated metadata columns.")
        currency = cells.text(row, 11)
        if currency != "MNT":
            raise _InvalidCell("unsupported_currency", "Only the verified MNT currency is supported.", "K")
        account = cells.text(row, 12)
        if re.fullmatch(r"IBAN: MN[0-9]+", account) is None:
            raise _InvalidCell("invalid_account", "Account identifier must be text in the verified IBAN field.", "L")
        start = _date(cells.text(row, 7), r"[0-9]{4}/[0-9]{2}/[0-9]{2}", "%Y/%m/%d", "G")[:10]
        end = _date(cells.text(row, 9), r"[0-9]{4}/[0-9]{2}/[0-9]{2}", "%Y/%m/%d", "I")[:10]
        printed = _date(cells.text(row, 2), r"[0-9]{2}/[0-9]{2}/[0-9]{4}", "%m/%d/%Y", "B")[:10]
        if start > end:
            raise _InvalidCell("invalid_period", "Statement period starts after it ends.")
        return StatementMetadata(parser_version=self.version, currency=currency,
                                 account_identifier=account.removeprefix("IBAN: "), worksheet=worksheet,
                                 period_start=start, period_end=end, printed_on=printed)

    @staticmethod
    def _totals(cells: _Cells, row: int) -> tuple[int, int]:
        for column in range(2, cells.original.max_column + 1):
            if column not in (4, 5) and not _blank(cells.value(row, column)):
                raise _InvalidCell("invalid_totals", "Totals row has unexpected populated columns.")
        credit, debit = cells.money(row, 4), cells.money(row, 5)
        if credit < 0 or debit < 0:
            raise _InvalidCell("invalid_totals", "Summary credit and debit totals must be nonnegative magnitudes.")
        return credit, debit

    @staticmethod
    def _transaction(cells: _Cells, row: int, metadata: StatementMetadata) -> ParsedTransaction:
        for column in range(9, cells.original.max_column + 1):
            if not _blank(cells.value(row, column)):
                raise _InvalidCell("unexpected_columns", "Transaction has populated columns outside the verified layout.")
        bank_datetime = cells.text(row, 1)
        occurred_at = _date(bank_datetime, _DATE_TIME, "%Y-%m-%d %H:%M:%S", "A")
        if not metadata.period_start <= occurred_at[:10] <= metadata.period_end:
            raise _InvalidCell("outside_period", "Transaction date falls outside the statement period.", "A")
        branch = cells.text(row, 2)
        if re.fullmatch(r"[0-9]+", branch) is None:
            raise _InvalidCell("invalid_branch", "Branch must be stored as digit text.", "B")
        opening, credit, debit, closing = (cells.money(row, c) for c in (3, 4, 5, 6))
        if credit < 0 or debit > 0:
            raise _InvalidCell("invalid_sign", "Credits must be nonnegative and row debits nonpositive.")
        if (credit != 0) == (debit != 0):
            raise _InvalidCell("ambiguous_amount", "Exactly one credit/debit amount must be nonzero.")
        description = cells.text(row, 7)
        counterparty = cells.text(row, 8, optional=True)
        if counterparty is not None and re.fullmatch(r"[A-Z]{0,2}[0-9]+", counterparty) is None:
            raise _InvalidCell("invalid_counterparty", "Counterparty account must be text in the observed account format.", "H")
        try:
            amount = require_integer(credit + debit, "Amount")
            source = ImportRow(description=description, bank_datetime=bank_datetime, occurred_at=occurred_at,
                               amount_minor=amount, currency=metadata.currency, balance_minor=closing,
                               worksheet=metadata.worksheet, row_number=row, reference=None)
        except ValueError:
            raise _InvalidCell("invalid_transaction", "Normalized transaction does not satisfy the money/date model.") from None
        return ParsedTransaction(source=source, branch=branch, counterparty_account=counterparty,
                                 opening_balance_minor=opening, credit_minor=credit, debit_minor=debit)
