"""Fictional Khan workbooks for the Step 7 manual and automated workflow."""

from pathlib import Path
from openpyxl import load_workbook
from tests.khan_fixture import create_fixture


def create_review_fixtures(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    merchant = "FABLE Жишээ CAFE* 1"
    cases = {
        "review-initial.xlsx": (10, [("A1B2C3 " + merchant, -1), ("D4E5F6 " + merchant, -1),
                                      ("G7H8J9 " + merchant, -1), ("QPay DEMO ONLY", -1),
                                      ("DEMO FAMILY CREDIT", 2)]),
        "review-future.xlsx": (11, [("K1L2M3 " + merchant, -1)]),
        "review-disabled.xlsx": (12, [("N4P5Q6 " + merchant, -1)]),
    }
    for filename, (day, rows) in cases.items():
        path = folder / filename
        create_fixture(path)
        book = load_workbook(path)
        try:
            sheet = book.active
            sheet.delete_rows(3, sheet.max_row)
            balance, credits, debits = 1000, 0, 0
            for index, (description, amount) in enumerate(rows):
                credit, debit = max(0, amount), min(0, amount)
                sheet.append([f"2026-10-{day:02d} 10:00:{index:02d}", "0001", balance,
                              credit, debit, balance + amount, description, None])
                balance += amount
                credits += credit
                debits -= debit
            sheet.append(["Нийт дүн:", None, None, f"{credits}.00", f"{debits}.00"])
            book.save(path)
        finally:
            book.close()
