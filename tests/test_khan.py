"""Synthetic workbook tests. Never locate or open private statements here."""

import hashlib
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from app.parsers.khan import HEADERS, KhanParser


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "khan_synthetic.xlsx"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class KhanParserTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory(prefix="parser-test-", dir=ROOT)
        self.addCleanup(temporary.cleanup)
        self.path = Path(temporary.name) / "invented.xlsx"
        shutil.copyfile(FIXTURE, self.path)
        self.parser = KhanParser()

    def change(self, action) -> None:
        book = load_workbook(self.path)
        try:
            action(book)
            book.save(self.path)
        finally:
            book.close()

    def cell(self, coordinate: str, value) -> None:
        self.change(lambda book: setattr(book.active[coordinate], "value", value))

    def xml_value(self, coordinate: str, value: str) -> None:
        """Inject a synthetic numeric lexical value or formula cache only."""
        with ZipFile(self.path) as archive:
            entries = {item.filename: archive.read(item.filename) for item in archive.infolist()}
        name = "xl/worksheets/sheet1.xml"
        root = ET.fromstring(entries[name])
        cell = next(c for c in root.iter(f"{{{NS}}}c") if c.get("r") == coordinate)
        element = cell.find(f"{{{NS}}}v")
        if element is None:
            element = ET.SubElement(cell, f"{{{NS}}}v")
        element.text = value
        entries[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        with ZipFile(self.path, "w") as archive:
            for entry, contents in entries.items():
                archive.writestr(entry, contents)

    def assert_issue(self, code: str, *, row: int | None = None):
        result = self.parser.parse(self.path)
        self.assertFalse(result.is_valid)
        found = [issue for issue in result.issues if issue.code == code]
        self.assertTrue(found, [issue.code for issue in result.issues])
        if row is not None:
            self.assertEqual(found[0].row_number, row)
        return result

    def test_fixture_credit_debit_dates_unicode_and_provenance(self) -> None:
        result = self.parser.parse(self.path)
        self.assertTrue(result.is_valid, result.issues)
        self.assertTrue(self.parser.detect(self.path))
        self.assertEqual(result.metadata.account_identifier, "MN00000000000000000000")
        self.assertEqual(result.metadata.currency, "MNT")
        self.assertEqual(result.metadata.period_start, "2026-10-01")
        self.assertEqual(result.metadata.period_end, "2026-10-31")
        self.assertEqual(result.metadata.printed_on, "2026-10-20")
        self.assertEqual([r.source.amount_minor for r in result.records], [10025, -2510, 15])
        self.assertEqual([r.source.row_number for r in result.records], [3, 4, 6])
        self.assertEqual(result.records[0].source.bank_datetime, "2026-10-05 09:10:11")
        self.assertEqual(result.records[0].source.occurred_at, "2026-10-05T09:10:11")
        self.assertEqual(result.records[0].source.description, "  Сургалтын дэмжлэг — туршилт  ")
        self.assertEqual(result.records[1].source.description, "Туршилтын кофе ☕")
        self.assertEqual(result.records[0].branch, "0001")
        self.assertEqual(result.records[0].counterparty_account, "00009999")
        self.assertIsNone(result.records[1].counterparty_account)
        self.assertTrue(all(r.source.reference is None for r in result.records))
        self.assertTrue(all(r.source.dedup_key is None for r in result.records))
        self.assertEqual(result.candidate_row_count, 3)
        self.assertTrue(all(c.outcome == "passed" for c in result.checks))

    def test_exact_observed_headers_are_in_the_fixture(self) -> None:
        book = load_workbook(FIXTURE)
        try:
            self.assertEqual(tuple(c.value for c in list(book.active.rows)[1][:8]), (
                "Гүйлгээний огноо", "Салбар", "Эхний үлдэгдэл", "Кредит гүйлгээ",
                "Дебит гүйлгээ", "Эцсийн үлдэгдэл", "Гүйлгээний утга", "Харьцсан данс",
            ))
        finally:
            book.close()

    def test_selection_uses_content_not_sheet_name_filename_or_header_row(self) -> None:
        def change(book):
            book.active.title = "Renamed synthetic sheet"
            book.active.insert_rows(1, 4)
            book.create_sheet("Unrelated notes")["A1"] = "Invented unrelated text"
        self.change(change)
        result = self.parser.parse(self.path)
        self.assertTrue(result.is_valid, result.issues)
        self.assertEqual([r.source.row_number for r in result.records], [7, 8, 10])
        self.assertEqual(result.records[0].source.worksheet, "Renamed synthetic sheet")

    def test_blank_rows_and_exact_repeated_headers_are_explicit(self) -> None:
        def change(book):
            sheet = book.active
            sheet.insert_rows(4)
            for column, label in enumerate(HEADERS, 1):
                sheet.cell(4, column, label)
            sheet.insert_rows(1)
        self.change(change)
        result = self.parser.parse(self.path)
        self.assertTrue(result.is_valid, result.issues)
        self.assertEqual(len(result.records), 3)
        self.assertEqual(result.issues[0].code, "repeated_header")
        self.assertEqual(result.issues[0].severity, "warning")
        self.assertEqual(result.issues[0].row_number, 5)

    def test_unsupported_headers_and_extra_header_columns_are_rejected(self) -> None:
        self.cell("E2", "Unverified debit label")
        self.assert_issue("unsupported_headers")
        self.assertFalse(self.parser.detect(self.path))
        self.cell("E2", HEADERS[4])
        self.cell("I2", "Unverified extra column")
        self.assert_issue("unsupported_headers")

    def test_unresolved_formula_in_extra_header_column_is_not_ignored(self) -> None:
        self.cell("I2", "=UNVERIFIED_COLUMN()")
        self.assert_issue("unsupported_headers")

    def test_unsupported_currency(self) -> None:
        self.cell("K1", "USD")
        self.assert_issue("unsupported_currency", row=1)

    def test_metadata_and_numeric_account_are_not_guessed(self) -> None:
        for coordinate, value, code in (("L1", 12345, "invalid_text"), ("L1", "IBAN: BAD", "invalid_account"),
                                        ("J1", "Unknown currency label", "invalid_metadata"),
                                        ("G1", "2026/11/01", "invalid_period")):
            with self.subTest(coordinate=coordinate):
                shutil.copyfile(FIXTURE, self.path)
                self.cell(coordinate, value)
                self.assert_issue(code, row=1)

    def test_multiple_matching_sheets_are_ambiguous(self) -> None:
        self.change(lambda book: book.copy_worksheet(book.active))
        self.assert_issue("ambiguous_worksheets")

    def test_unverified_merged_cells_are_rejected(self) -> None:
        self.change(lambda book: book.active.merge_cells("C1:D1"))
        self.assert_issue("unsupported_merged_cells")

    def test_unfamiliar_rows_are_not_silently_treated_as_footers(self) -> None:
        self.cell("A5", "Invented unknown row")
        result = self.assert_issue("invalid_date", row=5)
        self.assertEqual(result.candidate_row_count, 4)
        self.assertEqual(len(result.records), 3)

    def test_unknown_title_rows_are_rejected(self) -> None:
        def change(book):
            book.active.insert_rows(1)
            book.active["A1"] = "Unverified title row"
        self.change(change)
        self.assert_issue("unexpected_title_row", row=1)

    def test_content_after_totals_and_missing_totals_are_rejected(self) -> None:
        self.cell("A8", "Unverified footer")
        self.assert_issue("unexpected_after_totals", row=8)
        shutil.copyfile(FIXTURE, self.path)
        self.change(lambda book: book.active.delete_rows(7))
        self.assert_issue("missing_totals")

    def test_bad_money_precision_missing_values_and_locale_are_errors(self) -> None:
        for value in ("100.251", "1.234,56", "1,00.25", None, True, "NaN"):
            with self.subTest(value=value):
                shutil.copyfile(FIXTURE, self.path)
                self.cell("D3", value)
                result = self.assert_issue("invalid_money", row=3)
                self.assertEqual(result.issues[0].column, "D")

    def test_numeric_xml_precision_is_not_lost_through_float_conversion(self) -> None:
        self.xml_value("D3", "100.2500000000000001")
        self.assert_issue("invalid_money", row=3)

    def test_grouped_text_numbers_and_totals_use_verified_separators(self) -> None:
        self.cell("C3", "1,000.00")
        result = self.parser.parse(self.path)
        self.assertTrue(result.is_valid, result.issues)
        self.assertEqual(result.records[0].opening_balance_minor, 100000)

    def test_wrong_sign_and_both_or_neither_amounts_are_rejected(self) -> None:
        for credit, debit, code in ((-100.25, 0, "invalid_sign"), (100.25, 1, "invalid_sign"),
                                   (100.25, -1, "ambiguous_amount"), (0, 0, "ambiguous_amount")):
            with self.subTest(credit=credit, debit=debit):
                shutil.copyfile(FIXTURE, self.path)
                def change(book):
                    book.active["D3"], book.active["E3"] = credit, debit
                self.change(change)
                self.assert_issue(code, row=3)

    def test_bad_dates_are_not_interpreted_using_a_different_locale(self) -> None:
        for value in ("05/10/2026 09:10:11", "2026-02-30 09:10:11", "2026-10-05 25:00:00", "2026-10-05", 46000):
            with self.subTest(value=value):
                shutil.copyfile(FIXTURE, self.path)
                self.cell("A3", value)
                result = self.parser.parse(self.path)
                self.assertFalse(result.is_valid)
                self.assertEqual(result.issues[0].row_number, 3)
                self.assertIn(result.issues[0].code, ("invalid_date", "invalid_text"))

    def test_outside_period_date_is_reported(self) -> None:
        self.cell("A3", "2026-09-30 09:10:11")
        self.assert_issue("outside_period", row=3)

    def test_preserves_row_order_instead_of_sorting_bad_order(self) -> None:
        self.cell("A3", "2026-10-08 09:10:11")
        result = self.assert_issue("reconciliation_date_order")
        self.assertEqual([r.source.row_number for r in result.records], [3, 4, 6])

    def test_descriptions_branch_and_counterparty_are_not_coerced(self) -> None:
        for coordinate, value, code in (("G3", " ", "invalid_text"), ("B3", 1, "invalid_text"),
                                        ("H3", 9999, "invalid_text"), ("H3", "unverified account", "invalid_counterparty")):
            with self.subTest(coordinate=coordinate):
                shutil.copyfile(FIXTURE, self.path)
                self.cell(coordinate, value)
                self.assert_issue(code, row=3)

    def test_unexpected_transaction_columns(self) -> None:
        self.cell("I3", "Unverified extra field")
        self.assert_issue("unexpected_columns", row=3)

    def test_running_balance_and_row_balance_mismatches_are_errors(self) -> None:
        self.cell("C4", 1101.25)
        result = self.assert_issue("reconciliation_running_balances", row=4)
        self.assertIn("reconciliation_row_balances", [issue.code for issue in result.issues])
        self.assertTrue(all(issue.row_number == 4 for issue in result.issues))

    def test_totals_are_verified_and_debit_total_is_positive_magnitude(self) -> None:
        self.cell("E7", "-25.10")
        self.assert_issue("invalid_totals", row=7)
        self.cell("E7", "25.11")
        self.assert_issue("reconciliation_debit_total")
        self.cell("E7", "25.10")
        self.cell("D7", "100.41")
        self.assert_issue("reconciliation_credit_total")

    def test_formula_without_cache_is_rejected_without_computation(self) -> None:
        self.cell("D3", "=100.25")
        self.assert_issue("missing_formula_cache", row=3)

    def test_formula_with_usable_cache_is_read_without_evaluation(self) -> None:
        self.cell("D3", "=THIS_FORMULA_IS_NOT_EVALUATED()")
        self.xml_value("D3", "100.25")
        result = self.parser.parse(self.path)
        self.assertTrue(result.is_valid, result.issues)
        self.assertEqual(result.records[0].credit_minor, 10025)

    def test_formula_error_and_optional_formula_without_cache_are_not_omissions(self) -> None:
        self.cell("D3", "#VALUE!")
        self.assert_issue("cell_error", row=3)
        shutil.copyfile(FIXTURE, self.path)
        self.cell("H3", '=SYNTHETIC_ACCOUNT()')
        self.assert_issue("missing_formula_cache", row=3)

    def test_formula_only_row_is_not_treated_as_blank(self) -> None:
        self.cell("A5", "=NOW()")
        self.assert_issue("missing_formula_cache", row=5)

    def test_invalid_file_missing_file_and_wrong_extension_are_structured(self) -> None:
        self.assertEqual(self.parser.parse(self.path.with_suffix('.xls')).issues[0].code, "unsupported_file")
        missing = self.path.parent / "not-present.xlsx"
        self.assertEqual(self.parser.parse(missing).issues[0].code, "unreadable_workbook")
        self.path.write_bytes(b"synthetic invalid zip")
        self.assert_issue("unreadable_workbook")

    def test_malformed_workbook_library_errors_do_not_expose_source_values(self) -> None:
        with patch("app.parsers.khan.load_workbook", side_effect=TypeError("SYNTHETIC_PRIVATE_VALUE")):
            result = self.assert_issue("unreadable_workbook")
        self.assertNotIn("SYNTHETIC_PRIVATE_VALUE", repr(result))

    def test_workbooks_close_on_success_and_row_failure_and_open_failure(self) -> None:
        opened = []
        def tracked(*args, **kwargs):
            book = load_workbook(*args, **kwargs)
            book.close = Mock(wraps=book.close)
            opened.append(book)
            return book
        for malformed in (False, True):
            if malformed:
                self.cell("D3", "bad")
            opened.clear()
            with patch("app.parsers.khan.load_workbook", side_effect=tracked):
                self.parser.parse(self.path)
            self.assertEqual(len(opened), 2)
            self.assertTrue(all(book.close.call_count == 1 for book in opened))
        opened.clear()
        def fail_second(*args, **kwargs):
            if opened:
                raise ValueError("Synthetic cached-workbook failure")
            return tracked(*args, **kwargs)
        with patch("app.parsers.khan.load_workbook", side_effect=fail_second):
            self.assert_issue("unreadable_workbook")
        self.assertEqual(opened[0].close.call_count, 1)

    def test_diagnostics_and_result_repr_do_not_include_private_cell_values(self) -> None:
        secret = "SYNTHETIC_PRIVATE_VALUE_DO_NOT_LOG"
        self.cell("D3", secret)
        result = self.parser.parse(self.path)
        self.assertNotIn(secret, repr(result))
        self.assertTrue(all(secret not in issue.message for issue in result.issues))
        self.assertNotIn("MN00000000000000000000", repr(result))
        self.assertNotIn("Сургалтын", repr(result))

    def test_parser_never_writes_to_source(self) -> None:
        before = hashlib.sha256(self.path.read_bytes()).digest()
        self.parser.parse(self.path)
        self.parser.detect(self.path)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)


if __name__ == "__main__":
    unittest.main()
