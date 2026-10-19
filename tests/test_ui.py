"""Synthetic UI/service boundary checks. Hidden Tk checks are not visual QA."""

import os
import sqlite3
import tempfile
import threading
import time
import tkinter as tk
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from app.db import Repository
from app.import_service import ImportService
from app.models import ImportInfo, ImportRow
from app.parsers.base import ParseIssue
from app.ui.desktop import DesktopApp
from app.ui.presentation import browser_filters, format_mnt, issue_location, review_status
from app.ui.worker import FinanceWorker
from tests.khan_fixture import create_fixture
from tests.ui_fixtures import create_ui_fixtures


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "khan_synthetic.xlsx"


def temporary_case(case):
    base = ROOT / "data" / "test-temp"
    base.mkdir(parents=True, exist_ok=True)
    temporary = tempfile.TemporaryDirectory(prefix="ui-test-", dir=base)
    case.addCleanup(temporary.cleanup)
    return Path(temporary.name)


def synthetic_rows(count=23):
    return [ImportRow(description=f"Жишээ Cafe\u0301 %_ ' {i}", bank_datetime="2026-10-10 10:00:00",
                      occurred_at="2026-10-10T10:00:00", amount_minor=-100 if i % 2 else 100,
                      worksheet="Synthetic", row_number=i + 3) for i in range(count)]


class BrowserRepositoryTests(unittest.TestCase):
    def setUp(self):
        self.folder = temporary_case(self)
        self.repo = Repository(self.folder / "test.sqlite3")
        self.addCleanup(self.repo.close)
        self.account = self.repo.create_account("000123", "Synthetic")
        self.import_id = self.repo.write_import(
            ImportInfo(account_id=self.account, sha256="a" * 64, parser_version="synthetic", filename="synthetic.xlsx"),
            synthetic_rows())

    def test_direction_counts_and_pagination_are_consistent(self):
        filters = dict(direction="debit", review_only=True, account_id=self.account, date_from="2026-10-10", date_to="2026-10-10")
        total = self.repo.count_transactions(**filters)
        self.assertEqual(total, 11)
        pages = [self.repo.query_transactions(**filters, offset=i, limit=4) for i in range(0, total, 4)]
        identifiers = [row["id"] for page in pages for row in page]
        self.assertEqual(len(set(identifiers)), total)
        self.assertEqual(identifiers, sorted(identifiers, reverse=True))
        self.assertEqual(self.repo.count_transactions(direction="credit"), 12)
        self.assertEqual(self.repo.count_transactions(direction="zero"), 0)

    def test_unicode_search_is_literal_casefolded_and_normalized(self):
        for text in ("ЖИШЭЭ", "CAFÉ", "%_ '"):
            with self.subTest(text=text):
                self.assertEqual(self.repo.count_transactions(search=text), 23)
                self.assertEqual(len(self.repo.query_transactions(search=text, limit=5)), 5)
        self.assertEqual(self.repo.count_transactions(search="' OR 1=1 --"), 0)
        self.assertEqual(self.repo.count_transactions(search="%missing%"), 0)

    def test_category_and_review_filters_use_effective_values(self):
        category = next(c["id"] for c in self.repo.list_categories() if c["name"] == "Other expenses")
        row = self.repo.query_transactions(direction="debit", limit=1)[0]
        self.repo.assign_category(row["id"], category, reason="Synthetic classification")
        self.repo.edit_transaction(row["id"], description="Invented edited text", reason="Synthetic edit")
        self.assertEqual(self.repo.count_transactions(category_id=category, direction="debit"), 1)
        self.assertEqual(self.repo.count_transactions(review_only=True), 22)
        self.assertEqual(self.repo.count_transactions(search="EDITED"), 1)
        self.assertEqual(self.repo.count_transactions(category_id=category, review_only=True), 0)

    def test_invalid_filters_rejected_by_query_and_count(self):
        for kwargs in ({"direction": "income"}, {"date_from": "2026-02-30"},
                       {"date_from": "2026-10-11", "date_to": "2026-10-10"}, {"search": None}):
            for method in (self.repo.query_transactions, self.repo.count_transactions):
                with self.subTest(kwargs=kwargs, method=method.__name__), self.assertRaises(ValueError):
                    method(**kwargs)


class PresentationTests(unittest.TestCase):
    def test_entry_point_shuts_down_worker_when_event_loop_exits(self):
        from app.__main__ import main
        window = Mock()
        window.mainloop.side_effect = RuntimeError("Synthetic event loop exit")
        with patch("app.__main__.get_data_dir", return_value=ROOT / "data" / "synthetic"), patch("app.__main__.create_window", return_value=window):
            with self.assertRaises(RuntimeError):
                main()
        window.finance_app.worker.stop.assert_called_once_with()
        window.finance_app.worker.thread.join.assert_called_once_with()

    def test_money_format_is_exact_and_unicode_safe(self):
        self.assertEqual(format_mnt(-125050), "-1,250.50 MNT")
        self.assertEqual(format_mnt(0), "0.00 MNT")
        self.assertEqual(format_mnt(9223372036854775807), "92,233,720,368,547,758.07 MNT")
        self.assertEqual(format_mnt(None), "—")

    def test_dates_and_structured_error_coordinates(self):
        self.assertEqual(browser_filters(" 2026-10-10 ")["date_from"], "2026-10-10")
        for first, last in (("invalid", ""), ("2026-10-10T01:02:03", ""), ("2026-10-11", "2026-10-10")):
            with self.assertRaises(ValueError):
                browser_filters(first, last)
        self.assertEqual(issue_location(ParseIssue("bad", "Synthetic", sheet_index=0, row_number=7, column="E")), ("Worksheet 1", "7", "E"))
        self.assertEqual(issue_location(ParseIssue("bad", "Synthetic")), ("Statement", "—", "—"))

    def test_review_status_respects_explicit_other_and_excluded(self):
        self.assertEqual(review_status(dict(kind="unclassified", category_id=None, manual_lock=False)), "Needs review")
        self.assertEqual(review_status(dict(kind="expense", category_id=13, manual_lock=True)), "Manual / protected")
        self.assertEqual(review_status(dict(kind="excluded", category_id=None, manual_lock=False)), "Excluded")


class WorkerTests(unittest.TestCase):
    def setUp(self):
        self.folder = temporary_case(self)
        self.database = self.folder / "finance.sqlite3"

    def start(self, **options):
        self.worker = FinanceWorker(self.database, **options)
        self.addCleanup(self.stop)
        event = self.worker.events.get(timeout=5)
        self.assertEqual(event.operation, "ready")
        self.account = self.ask("account", name="Synthetic Khan", identifier="MN00000000000000000000").payload["selected"]

    def stop(self):
        self.worker.stop()
        self.worker.thread.join(5)
        self.assertFalse(self.worker.thread.is_alive())

    def ask(self, operation, **payload):
        self.worker.submit(operation, **payload)
        event = self.worker.events.get(timeout=5)
        self.assertEqual(event.operation, operation)
        return event

    def snapshot(self):
        with closing(sqlite3.connect(self.database)) as connection:
            return tuple(connection.iterdump())

    def test_preview_cancel_is_read_only_then_commit_and_reopen(self):
        self.start()
        before = self.snapshot()
        event = self.ask("preview", path=FIXTURE, account_id=self.account)
        self.assertTrue(event.payload.is_valid)
        self.assertEqual(self.snapshot(), before)
        self.assertIsNone(self.ask("cancel").error)
        self.assertEqual(self.snapshot(), before)
        self.ask("preview", path=FIXTURE, account_id=self.account)
        result = self.ask("commit", decisions={})
        self.assertIsNone(result.error)
        self.assertEqual(result.payload["result"].imported_count, 3)
        self.assertEqual(result.payload["unclassified"], 3)
        self.ask("preview", path=FIXTURE, account_id=self.account)
        identical = self.ask("commit", decisions={}).payload
        self.assertEqual(identical["result"].status, "identical_file")
        self.assertEqual(identical["unclassified"], 0)
        self.stop()
        with Repository(self.database) as reopened:
            self.assertEqual(reopened.count_transactions(), 3)

    def test_database_and_services_live_in_one_worker_thread(self):
        owners = []
        class TrackingRepository(Repository):
            def __init__(self, path):
                owners.append(threading.get_ident())
                super().__init__(path)
            def close(self):
                owners.append(threading.get_ident())
                super().close()
        class TrackingService(ImportService):
            def preview(self, *args):
                owners.append(threading.get_ident())
                return super().preview(*args)
            def confirm(self, *args):
                owners.append(threading.get_ident())
                return super().confirm(*args)
        self.start(repository_factory=TrackingRepository, service_factory=TrackingService)
        self.ask("preview", path=FIXTURE, account_id=self.account)
        self.ask("commit", decisions={})
        self.stop()
        self.assertEqual(set(owners), {self.worker.thread.ident})
        self.assertNotIn(threading.get_ident(), owners)

    def test_malformed_preview_cannot_commit_and_worker_recovers(self):
        self.start()
        bad = self.folder / "bad.xlsx"
        bad.write_bytes(b"FICTIONAL INVALID WORKBOOK")
        before = self.snapshot()
        preview = self.ask("preview", path=bad, account_id=self.account).payload
        self.assertFalse(preview.is_valid)
        self.assertIsNotNone(self.ask("commit", decisions={}).error)
        self.assertEqual(self.snapshot(), before)
        self.ask("cancel")
        self.assertTrue(self.ask("preview", path=FIXTURE, account_id=self.account).payload.is_valid)

    def test_overlap_needs_explicit_decisions_and_records_them(self):
        self.start()
        self.ask("preview", path=FIXTURE, account_id=self.account)
        self.ask("commit", decisions={})
        overlap = self.folder / "overlap.xlsx"
        book = load_workbook(FIXTURE)
        try:
            book.active.title = "Synthetic overlap"
            book.save(overlap)
        finally:
            book.close()
        preview = self.ask("preview", path=overlap, account_id=self.account).payload
        self.assertEqual(preview.ambiguous_count, 3)
        before = self.snapshot()
        self.assertIsNotNone(self.ask("commit", decisions={}).error)
        self.assertEqual(self.snapshot(), before)
        self.assertIsNone(self.ask("commit", decisions={r.index: "skip" for r in preview.rows}).error)
        with Repository(self.database) as repo:
            self.assertEqual(repo.count_transactions(), 3)
            self.assertTrue(all(d["action"] == "skip" for d in repo.get_import_decisions(2)))

    def test_manual_check_fixtures_support_paging_overlap_and_row_errors(self):
        self.start()
        fixtures = self.folder / "manual-fixtures"
        create_ui_fixtures(fixtures)
        preview = self.ask("preview", path=fixtures / "synthetic-browser.xlsx", account_id=self.account).payload
        self.assertTrue(preview.is_valid)
        self.assertEqual(preview.valid_count, 23)
        self.ask("commit", decisions={})
        pages = [self.ask("browse", page=page, page_size=10).payload for page in range(3)]
        self.assertEqual([len(page["rows"]) for page in pages], [10, 10, 3])
        self.assertTrue(all(page["total"] == 23 for page in pages))
        overlap = self.ask("preview", path=fixtures / "synthetic-overlap.xlsx", account_id=self.account).payload
        self.assertEqual(overlap.ambiguous_count, 23)
        self.ask("cancel")
        bad = self.ask("preview", path=fixtures / "synthetic-invalid.xlsx", account_id=self.account).payload
        self.assertFalse(bad.is_valid)
        self.assertTrue(any(i.sheet_index == 0 and i.row_number == 4 for i in bad.issues))

    def test_write_failure_rolls_back_and_keeps_preview_cancelable(self):
        class FailingRepository(Repository):
            def _apply_classification(self, *args):
                super()._apply_classification(*args)
                raise sqlite3.IntegrityError("SYNTHETIC_PRIVATE_FAILURE")
        self.start(repository_factory=FailingRepository)
        fixture = self.folder / "purchase.xlsx"
        book = load_workbook(FIXTURE)
        try:
            book.active["G4"] = "A1B2C3 DEMO STORE"
            book.save(fixture)
        finally:
            book.close()
        self.ask("preview", path=fixture, account_id=self.account)
        before = self.snapshot()
        result = self.ask("commit", decisions={})
        self.assertIsNotNone(result.error)
        self.assertNotIn("SYNTHETIC_PRIVATE_FAILURE", result.error)
        self.assertEqual(self.snapshot(), before)
        self.assertIsNone(self.ask("cancel").error)

    def test_post_commit_count_error_still_reports_committed(self):
        class CountFailureRepository(Repository):
            def count_transactions(self, **kwargs):
                if kwargs.get("import_id"):
                    raise sqlite3.OperationalError("SYNTHETIC_PRIVATE_FAILURE")
                return super().count_transactions(**kwargs)
        self.start(repository_factory=CountFailureRepository)
        self.ask("preview", path=FIXTURE, account_id=self.account)
        result = self.ask("commit", decisions={})
        self.assertIsNone(result.error)
        self.assertEqual(result.payload["result"].status, "committed")
        self.assertIsNone(result.payload["unclassified"])
        self.assertEqual(self.ask("browse").payload["total"], 3)

    def test_worker_errors_do_not_expose_exception_text_and_can_recover(self):
        class BadService(ImportService):
            def preview(self, *args):
                raise ValueError("SYNTHETIC_PRIVATE_TEXT")
        self.start(service_factory=BadService)
        result = self.ask("preview", path=FIXTURE, account_id=self.account)
        self.assertNotIn("SYNTHETIC_PRIVATE_TEXT", repr(result))
        self.assertIsNotNone(result.error)
        self.assertEqual(self.ask("browse").payload["total"], 0)

    def test_shutdown_during_parse_waits_and_does_not_import(self):
        entered, release = threading.Event(), threading.Event()
        class SlowService(ImportService):
            def preview(self, *args):
                entered.set()
                release.wait(5)
                return super().preview(*args)
        self.start(service_factory=SlowService)
        self.addCleanup(release.set)
        self.worker.submit("preview", path=FIXTURE, account_id=self.account)
        self.assertTrue(entered.wait(3))
        self.worker.stop()
        self.assertTrue(self.worker.thread.is_alive())
        release.set()
        self.worker.thread.join(5)
        self.assertFalse(self.worker.thread.is_alive())
        with Repository(self.database) as repo:
            self.assertEqual(repo.count_transactions(), 0)

    def test_startup_failure_is_safe_and_stops(self):
        def fail(_):
            raise OSError("SYNTHETIC_PRIVATE_PATH")
        self.worker = FinanceWorker(self.database, repository_factory=fail)
        self.addCleanup(self.stop)
        fatal = self.worker.events.get(timeout=5)
        self.assertEqual(fatal.operation, "fatal")
        self.assertNotIn("SYNTHETIC_PRIVATE_PATH", fatal.error)
        self.assertEqual(self.worker.events.get(timeout=5).operation, "stopped")


class HiddenWidgetTests(unittest.TestCase):
    """Construct withdrawn Tk widgets; never move a pointer or generate key input."""

    def setUp(self):
        if os.environ.get("PERSONAL_FINANCE_SKIP_GUI_TESTS") == "1":
            self.skipTest("GUI tests explicitly disabled; manual GUI acceptance remains pending")
        self.folder = temporary_case(self)
        try:
            self.root = tk.Tk()
        except tk.TclError:
            self.skipTest("Tk could not initialize in this execution environment; GUI checks remain NOT RUN")
        self.root.withdraw()
        self.app = None
        self.addCleanup(self.clean_up)
        self.errors = []
        self.patch = patch("app.ui.desktop.messagebox.showerror", side_effect=lambda *a, **k: self.errors.append(a))
        self.patch.start()
        self.addCleanup(self.patch.stop)

    def clean_up(self):
        if self.app is not None:
            self.app.worker.stop()
            self.app.worker.thread.join(5)
            self.assertFalse(self.app.worker.thread.is_alive())
            try:
                self.root.after_cancel(self.app._poll_id)
            except tk.TclError:
                pass
        try:
            self.root.destroy()
        except tk.TclError:
            pass

    def pump(self, predicate):
        deadline = time.monotonic() + 6
        while time.monotonic() < deadline:
            self.root.update()
            if predicate():
                return
            time.sleep(0.01)
        self.fail("Synthetic UI operation did not finish in time")

    def start(self, service_factory=ImportService):
        self.app = DesktopApp(self.root, self.folder,
                              worker_factory=lambda path: FinanceWorker(path, service_factory=service_factory))
        self.pump(lambda: self.app.ready and not self.app.busy)
        self.app._request("account", name="Synthetic Khan", identifier="MN00000000000000000000")
        self.pump(lambda: not self.app.busy)

    def preview(self, path=FIXTURE):
        with patch("app.ui.desktop.filedialog.askopenfilename", return_value=str(path)):
            self.app.choose_button.invoke()
        self.assertTrue(self.app.preview_button.instate(["disabled"]))
        self.app.account_confirmed.set(True)
        self.app._sync_controls()
        self.app.preview_button.invoke()
        self.assertTrue(self.app.busy)
        self.assertTrue(self.app.choose_button.instate(["disabled"]))
        self.pump(lambda: not self.app.busy)

    def test_cancel_import_browse_unicode_filter_and_reopen(self):
        self.start()
        self.preview()
        self.app.cancel_button.invoke()
        self.pump(lambda: not self.app.busy)
        with Repository(self.folder / "finance.sqlite3") as repo:
            self.assertEqual(repo.count_transactions(), 0)
        self.preview()
        self.app.commit_button.invoke()
        self.pump(lambda: not self.app.busy)
        self.assertIn("Added: 3", self.app.summary.get())
        self.assertIn("Unclassified added: 3", self.app.summary.get())
        self.assertEqual(len(self.app.transactions.tree.get_children()), 3)
        self.app.search.set("КОФЕ")
        self.app.direction_text.set("Debit")
        self.app._apply_filters()
        self.pump(lambda: not self.app.busy)
        self.assertEqual(len(self.app.transactions.tree.get_children()), 1)
        self.app.review_button.invoke()
        self.pump(lambda: not self.app.busy)
        self.assertTrue(self.app.review_only.get())
        self.assertEqual(len(self.app.transactions.tree.get_children()), 3)
        self.assertFalse(self.errors)
        self.app.close()
        self.pump(lambda: not self.app.worker.thread.is_alive())
        with Repository(self.folder / "finance.sqlite3") as repo:
            self.assertEqual(repo.count_transactions(), 3)

    def test_pagination_category_filter_and_invalid_dates(self):
        self.start()
        with Repository(self.folder / "finance.sqlite3") as repo:
            account = repo.list_accounts()[0]["id"]
            repo.write_import(ImportInfo(account_id=account, sha256="c" * 64, parser_version="synthetic", filename="synthetic.xlsx"), synthetic_rows())
            category = next(c["id"] for c in repo.list_categories() if c["name"] == "Other expenses")
            row = repo.query_transactions(direction="debit", limit=1)[0]
            repo.assign_category(row["id"], category, reason="Synthetic")
        self.app.page_size.set("10")
        self.app._browse(0)
        self.pump(lambda: not self.app.busy)
        self.assertEqual(len(self.app.transactions.tree.get_children()), 10)
        self.app.browser_next.invoke()
        self.pump(lambda: not self.app.busy)
        self.assertEqual(self.app.browser_page, 1)
        self.app.browser_next.invoke()
        self.pump(lambda: not self.app.busy)
        self.assertEqual(len(self.app.transactions.tree.get_children()), 3)
        self.assertTrue(self.app.browser_next.instate(["disabled"]))
        self.app.category_text.set("Other expenses")
        self.app._apply_filters()
        self.pump(lambda: not self.app.busy)
        self.assertEqual(self.app.browser_total, 1)
        self.app.date_from.set("invalid synthetic date")
        self.app._apply_filters()
        self.assertEqual(len(self.errors), 1)
        self.assertFalse(self.app.busy)
        self.assertEqual(self.app.browser_total, 1)

    def test_ambiguous_decisions_are_explicit_and_candidates_visible(self):
        self.start()
        self.preview()
        self.app.commit_button.invoke()
        self.pump(lambda: not self.app.busy)
        other = self.folder / "overlap.xlsx"
        book = load_workbook(FIXTURE)
        try:
            book.active.title = "Synthetic overlap"
            book.save(other)
        finally:
            book.close()
        self.preview(other)
        self.assertEqual(self.app.decisions, {})
        self.assertTrue(self.app.commit_button.instate(["disabled"]))
        self.app.preview_table.tree.selection_set("0")
        self.root.update()
        self.assertEqual(len(self.app.candidates.items), 2)
        self.app.decision_buttons[1].invoke()
        self.assertEqual(self.app.decisions, {0: "skip"})
        self.assertTrue(self.app.commit_button.instate(["disabled"]))
        self.app.preview_table.tree.selection_set("1", "2")
        self.root.update()
        self.app.decision_buttons[1].invoke()
        self.assertTrue(self.app.commit_button.instate(["!disabled"]))
        self.app.commit_button.invoke()
        self.pump(lambda: not self.app.busy)
        self.assertIn("Skipped: 3", self.app.summary.get())
        self.assertEqual(self.app.browser_total, 3)
        self.assertFalse(self.errors)

    def test_invalid_rows_show_sheet_and_row_and_block_commit(self):
        self.start()
        bad = self.folder / "bad-row.xlsx"
        create_fixture(bad)
        book = load_workbook(bad)
        try:
            book.active["E4"] = "not synthetic money"
            book.save(bad)
        finally:
            book.close()
        self.preview(bad)
        self.assertTrue(self.app.commit_button.instate(["disabled"]))
        values = [self.app.issues.tree.item(i, "values") for i in self.app.issues.tree.get_children()]
        self.assertTrue(any(row[1] == "Worksheet 1" and row[2] == "4" for row in values))
        self.app.cancel_button.invoke()
        self.pump(lambda: not self.app.busy)
        self.preview()
        self.assertTrue(self.app.preview.is_valid)

    def test_tk_event_loop_runs_while_parsing_and_close_waits(self):
        entered, release = threading.Event(), threading.Event()
        class SlowService(ImportService):
            def preview(self, *args):
                entered.set()
                release.wait(5)
                return super().preview(*args)
        self.addCleanup(release.set)
        self.start(service_factory=SlowService)
        self.app.file_path.set(str(FIXTURE))
        self.app.account_confirmed.set(True)
        self.app._preview()
        self.assertTrue(entered.wait(3))
        ticks = []
        self.root.after(0, lambda: ticks.append(True))
        self.root.update()
        self.assertEqual(ticks, [True])
        self.assertTrue(self.app.commit_button.instate(["disabled"]))
        self.assertTrue(self.app.choose_button.instate(["disabled"]))
        self.app.close()
        release.set()
        self.pump(lambda: not self.app.worker.thread.is_alive())
        with Repository(self.folder / "finance.sqlite3") as repo:
            self.assertEqual(repo.count_transactions(), 0)


if __name__ == "__main__":
    unittest.main()
