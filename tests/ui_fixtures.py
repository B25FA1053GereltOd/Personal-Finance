"""Generate fictional workbooks for manual UI checks inside a chosen folder."""

from pathlib import Path

from openpyxl import load_workbook

from tests.khan_fixture import create_fixture


def create_ui_fixtures(folder: Path) -> None:
    """23 rows, an overlapping copy, and an invalid row; never read real files."""
    folder.mkdir(parents=True, exist_ok=True)
    main = folder / "synthetic-browser.xlsx"
    create_fixture(main)
    book = load_workbook(main)
    try:
        sheet = book.active
        sheet.delete_rows(3, sheet.max_row)
        for index in range(23):
            sheet.append([f"2026-10-10 10:00:{index:02d}", "0001", 1000 - index, 0, -1, 999 - index,
                          f"A1B2C3 DEMO Жишээ STORE {index}", None])
        sheet.append(["Нийт дүн:", None, None, "0.00", "23.00"])
        book.save(main)
        sheet.title = "Synthetic overlap"
        book.save(folder / "synthetic-overlap.xlsx")
        sheet["E4"] = "invalid synthetic money"
        book.save(folder / "synthetic-invalid.xlsx")
    finally:
        book.close()
