"""Classification acceptance using fictional values only; no real statements."""

import sqlite3
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app import db
from app.classification import ClassificationBlocked, ClassificationService, extract_merchant, normalize_merchant_key
from app.db import Repository
from app.import_service import ImportBlocked, ImportService
from app.models import ImportInfo, ImportRow, MerchantMatch
from app.parsers.base import ParseResult, ParsedTransaction, StatementMetadata


def synthetic_extractor(bank, parser_version, description):
    """Test-only evidence. These delimiters are not a real bank format."""
    if bank not in ("Synthetic Bank A", "Synthetic Bank B") or parser_version != "synthetic-pos-v1":
        return None
    for prefix, kind in (("TEST PURCHASE|", "purchase"), ("TEST REFUND|", "refund")):
        if description.startswith(prefix) and description[len(prefix):].strip():
            return MerchantMatch(display_text=description[len(prefix):], pattern="synthetic-pos-v1", transaction_type=kind)
    return None


class SyntheticParser:
    version = "synthetic-pos-v1"

    def parse_bytes(self, payload):
        return self.result


class ClassificationTests(unittest.TestCase):
    def setUp(self):
        root = Path(__file__).resolve().parents[1]
        temp_root = root / "data" / "test-temp"
        temp_root.mkdir(parents=True, exist_ok=True)
        temporary = tempfile.TemporaryDirectory(prefix="classification-test-", dir=temp_root)
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        self.path = self.folder / "synthetic.sqlite3"
        self.repo = Repository(self.path)
        self.addCleanup(self.repo.close)
        self.accounts = [self.repo.create_account("000101", "Synthetic A", bank="Synthetic Bank A"),
                         self.repo.create_account("000102", "Synthetic B", bank="Synthetic Bank B")]
        self.expenses = [c["id"] for c in self.repo.list_categories() if c["direction"] == "expense"]
        self.income = next(c["id"] for c in self.repo.list_categories() if c["direction"] == "income")
        self.other = next(c["id"] for c in self.repo.list_categories() if c["name"] == "Other expenses")
        self.classifier = ClassificationService(self.repo, extractor=synthetic_extractor)
        self.parser = SyntheticParser()
        self.importer = ImportService(self.repo, self.parser, classification=self.classifier)
        self.sequence = 0

    def prepare(self, descriptions, *, amounts=None, account=None):
        self.sequence += 1
        account = self.accounts[0] if account is None else account
        identifier = next(a["identifier"] for a in self.repo.list_accounts() if a["id"] == account)
        amounts = [-100] * len(descriptions) if amounts is None else amounts
        sources = [ImportRow(description=description, bank_datetime=f"2026-10-10 10:00:{i:02d}",
                             occurred_at=f"2026-10-10T10:00:{i:02d}", amount_minor=amount,
                             worksheet="Synthetic", row_number=i + 3)
                   for i, (description, amount) in enumerate(zip(descriptions, amounts))]
        self.parser.result = ParseResult(
            metadata=StatementMetadata(parser_version=self.parser.version, currency="MNT", account_identifier=identifier,
                                       worksheet="Synthetic", period_start="2026-10-01", period_end="2026-10-31", printed_on="2026-11-01"),
            records=tuple(ParsedTransaction(source=source, branch="0001", counterparty_account=None,
                                            opening_balance_minor=0, credit_minor=max(source.amount_minor, 0),
                                            debit_minor=min(source.amount_minor, 0)) for source in sources),
            candidate_row_count=len(sources),
        )
        path = self.folder / f"synthetic-{self.sequence}.xlsx"
        path.write_bytes(f"INVENTED POS TEST FILE {self.sequence}".encode())
        return self.importer.preview(path, account)

    def import_rows(self, descriptions, **kwargs):
        preview = self.prepare(descriptions, **kwargs)
        # This helper explicitly keeps synthetic repeated occurrences; it does
        # not bypass the real import service's ambiguous-overlap confirmation.
        choices = {row.index: "keep" for row in preview.rows if row.status == "ambiguous"}
        result = self.importer.confirm(preview, choices)
        return [self.repo.get_transaction(row["id"]) for row in
                sorted(self.repo.query_transactions(import_id=result.import_id), key=lambda r: r["id"])]

    def rule(self, merchant="Store 12", category=None, bank="Synthetic Bank A", direction="debit"):
        category = self.expenses[0] if category is None else category
        return self.classifier.save_rule(bank, merchant, direction, category)

    def counts(self):
        return tuple(self.repo._connection.execute(f"SELECT count(*) FROM {name}").fetchone()[0]
                     for name in ("imports", "transactions", "import_decisions", "transaction_edits"))

    def test_repeated_merchant_matches_rule_and_records_id(self):
        rule = self.rule()
        rows = self.import_rows(["TEST PURCHASE|Store 12", "TEST PURCHASE|Store 12"])
        for row in rows:
            self.assertEqual(row["category_id"], self.expenses[0])
            self.assertEqual(row["classification_source"], "rule")
            self.assertEqual(row["applied_rule_id"], rule)
            self.assertEqual(row["applied_rule_revision"], 1)
            self.assertEqual(row["manual_lock"], 0)
            self.assertEqual(row["description"], row["original_description"])
            self.assertEqual(row["merchant_display"], "Store 12")
            self.assertEqual(self.repo.get_history(row["id"])[0]["changes"]["applied_rule_id"]["new"], rule)

    def test_unicode_casefold_and_whitespace_normalization_preserves_display(self):
        rule = self.rule(merchant="CAFÉ ЖИШЭЭ 12")
        display = "  Cafe\u0301\tЖишээ\u00a0\n12  "
        row = self.import_rows(["TEST PURCHASE|" + display])[0]
        self.assertEqual(row["merchant_key"], "café жишээ 12")
        self.assertEqual(row["merchant_display"], display)
        self.assertEqual(row["applied_rule_id"], rule)
        self.assertEqual(normalize_merchant_key("Straße"), normalize_merchant_key("STRASSE"))

    def test_digits_branches_punctuation_and_compatibility_characters_stay_distinct(self):
        names = ("Store 12", "Store 13", "Store 12 West", "Store-12", "Store.12", "Store １２")
        self.assertEqual(len({normalize_merchant_key(n) for n in names}), len(names))
        self.rule()
        rows = self.import_rows(["TEST PURCHASE|" + name for name in names])
        self.assertEqual([r["category_id"] for r in rows], [self.expenses[0], None, None, None, None, None])

    def test_exact_matches_only_not_substring_or_pattern_rules(self):
        self.rule(merchant="Store.*")
        rows = self.import_rows(["TEST PURCHASE|Store.*", "TEST PURCHASE|Store 12", "TEST PURCHASE|Big Store.*"])
        self.assertEqual([r["category_id"] for r in rows], [self.expenses[0], None, None])

    def test_bank_scope_and_one_rule_per_normalized_scope(self):
        a = self.rule()
        self.assertEqual(self.rule(merchant=" STORE\t12 ", category=self.expenses[1]), a)
        b = self.rule(bank="Synthetic Bank B", category=self.expenses[2])
        self.assertNotEqual(a, b)
        first = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        second = self.import_rows(["TEST PURCHASE|Store 12"], account=self.accounts[1])[0]
        self.assertEqual((first["applied_rule_id"], second["applied_rule_id"]), (a, b))
        self.assertEqual(len(self.repo.list_rules()), 2)

    def test_debit_rule_does_not_classify_refund_transfer_or_non_pos_credit(self):
        self.rule()
        rows = self.import_rows(["TEST PURCHASE|Store 12", "TEST REFUND|Store 12", "TEST TRANSFER|Store 12",
                                 "Synthetic unrelated credit", "TEST PURCHASE|Store 12"], amounts=[-100, 100, -100, 100, 100])
        self.assertEqual([r["classification_source"] for r in rows], ["rule"] + ["unclassified"] * 4)
        self.assertIsNone(rows[-1]["merchant_key"])
        self.assertIsNone(rows[2]["merchant_key"])

    def test_explicit_credit_rule_matches_only_verified_pos_refund(self):
        debit = self.rule()
        credit = self.rule(direction="credit", category=self.income)
        rows = self.import_rows(["TEST PURCHASE|Store 12", "TEST REFUND|Store 12", "Non-POS credit"], amounts=[-100, 100, 100])
        self.assertEqual([r["applied_rule_id"] for r in rows], [debit, credit, None])

    def test_missing_merchant_remains_individually_reviewable(self):
        rows = self.import_rows(["Unrecognized synthetic text", "TEST PURCHASE|   "])
        queue = self.classifier.review()
        self.assertEqual(queue.groups, ())
        self.assertEqual([item.transaction_id for item in queue.individual], [r["id"] for r in rows])
        self.assertTrue(all(r["category_id"] is None for r in rows))

    def test_default_khan_extractor_rejects_unverified_patterns(self):
        self.assertIsNone(extract_merchant("Khan Bank", "khan-mnt-1", "POS synthetic merchant / 12"))
        self.assertIsNone(extract_merchant("Khan Bank", "khan-mnt-1", "ПОС synthetic merchant"))
        self.importer = ImportService(self.repo, self.parser)
        self.rule()
        row = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        self.assertIsNone(row["merchant_key"])
        self.assertIsNone(row["category_id"])

    def test_real_parser_synthetic_fixture_without_card_prefix_stays_reviewable(self):
        account = self.repo.create_account("MN00000000000000000000", "Synthetic Khan")
        importer = ImportService(self.repo)
        fixture = Path(__file__).parent / "fixtures" / "khan_synthetic.xlsx"
        result = importer.confirm(importer.preview(fixture, account))
        rows = self.repo.query_transactions(import_id=result.import_id)
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(r["merchant_key"] is None and r["category_id"] is None for r in rows))
        self.assertEqual(len(ClassificationService(self.repo).review().individual), 3)

    def test_review_groups_exact_known_keys_by_bank_and_direction(self):
        first = self.import_rows(["TEST PURCHASE| Store 12 ", "TEST PURCHASE|STORE 12", "TEST REFUND|Store 12", "Unknown"],
                                 amounts=[-100, -100, 100, -100])
        self.import_rows(["TEST PURCHASE|Store 12"], account=self.accounts[1])
        queue = self.classifier.review()
        self.assertEqual(sorted(len(group.items) for group in queue.groups), [1, 1, 2])
        self.assertEqual([r.transaction_id for r in queue.individual], [first[3]["id"]])
        self.assertEqual(first[0]["merchant_display"], " Store 12 ")

    def test_transaction_only_assignment_creates_no_rule_and_locks_selection(self):
        rows = self.import_rows(["TEST PURCHASE|Store 12", "TEST PURCHASE|Store 12"])
        self.classifier.assign_category([rows[0]["id"]], self.expenses[0])
        first, second = [self.repo.get_transaction(r["id"]) for r in rows]
        self.assertEqual((first["category_id"], first["manual_lock"], first["classification_source"]),
                         (self.expenses[0], 1, "manual"))
        self.assertIsNone(second["category_id"])
        self.assertEqual(self.repo.list_rules(), [])

    def test_manually_assigned_other_is_classified_but_unknown_never_defaults(self):
        rows = self.import_rows(["Unknown synthetic item", "Unknown second item"])
        self.classifier.reclassify(rows[0]["id"], self.other)
        queue = self.classifier.review()
        self.assertEqual([item.transaction_id for item in queue.individual], [rows[1]["id"]])
        self.assertIsNone(self.repo.get_transaction(rows[1]["id"])["category_id"])

    def test_saving_rule_only_affects_future_imports(self):
        old = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        self.rule()
        self.assertIsNone(self.repo.get_transaction(old["id"])["category_id"])
        new = self.import_rows(["TEST PURCHASE|Store 12", "Unique unrelated row"])[0]
        self.assertEqual(new["category_id"], self.expenses[0])

    def test_import_callback_cannot_be_reused_to_reclassify_past_rows(self):
        row = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        self.rule()
        with self.assertRaisesRegex(ValueError, "preview a rule"):
            self.classifier.classify_import(row["import_id"])
        self.assertEqual(self.repo.get_transaction(row["id"]), row)

    def test_rule_edit_does_not_change_past_category_or_recorded_revision(self):
        rule = self.rule()
        old = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        self.assertEqual(self.rule(category=self.expenses[1]), rule)
        self.assertEqual(self.repo.get_transaction(old["id"]), old)
        self.assertEqual(self.repo.get_rule(rule)["revision"], 2)

    def test_manual_override_survives_rule_update_and_reimport(self):
        rule = self.rule()
        preview = self.prepare(["TEST PURCHASE|Store 12"])
        self.importer.confirm(preview)
        row = self.repo.query_transactions(account_id=self.accounts[0])[0]
        self.classifier.reclassify(row["id"], self.other)
        before = self.repo.get_transaction(row["id"])
        self.assertEqual(before["manual_lock"], 1)
        self.assertIsNone(before["applied_rule_id"])
        self.rule(category=self.expenses[1])
        overlap = self.prepare(["TEST PURCHASE|Store 12"])
        self.importer.confirm(overlap, {0: "skip"})
        self.assertEqual(self.repo.get_transaction(row["id"]), before)
        self.assertEqual(self.classifier.preview_rule(rule).candidates, ())

    def test_disabling_rule_preserves_history_and_stops_future_matches(self):
        rule = self.rule()
        old = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        self.classifier.disable_rule(rule)
        self.assertEqual(self.repo.get_transaction(old["id"]), old)
        self.assertEqual(self.repo.get_rule(rule)["active"], 0)
        self.assertEqual(self.repo.list_rules(), [])
        self.assertEqual(len(self.repo.list_rules(include_disabled=True)), 1)
        new = self.import_rows(["Some unrelated row", "TEST PURCHASE|Store 12"])[1]
        self.assertIsNone(new["category_id"])
        self.assertEqual(self.rule(), rule)
        self.assertEqual(self.repo.get_rule(rule)["active"], 1)
        self.assertIsNone(self.repo.get_transaction(new["id"])["category_id"])

    def test_retroactive_apply_changes_only_explicitly_selected_eligible_rows(self):
        rows = self.import_rows(["TEST PURCHASE|Store 12"] * 4 + ["TEST PURCHASE|Different", "TEST REFUND|Store 12"],
                                amounts=[-100] * 5 + [100])
        self.classifier.reclassify(rows[2]["id"], self.other)
        self.repo.edit_transaction(rows[3]["id"], reason="Explicit synthetic transfer", kind="transfer", manual_lock=False)
        other = self.import_rows(["TEST PURCHASE|Store 12"], account=self.accounts[1])[0]
        rule = self.rule()
        preview = self.classifier.preview_rule(rule)
        self.assertEqual([r.transaction_id for r in preview.candidates], [r["id"] for r in rows[:2]])
        before = {r["id"]: self.repo.get_transaction(r["id"]) for r in rows[1:] + [other]}
        self.assertEqual(self.classifier.apply_rule(preview, [rows[0]["id"]]), 1)
        self.assertEqual(self.repo.get_transaction(rows[0]["id"])["applied_rule_id"], rule)
        for identifier, original in before.items():
            self.assertEqual(self.repo.get_transaction(identifier), original)

    def test_preview_and_cancel_write_nothing(self):
        self.import_rows(["TEST PURCHASE|Store 12"])
        rule = self.rule()
        before = self.path.read_bytes()
        preview = self.classifier.preview_rule(rule)
        self.classifier.cancel_rule_preview(preview)
        self.assertEqual(self.path.read_bytes(), before)
        with self.assertRaises(ClassificationBlocked):
            self.classifier.apply_rule(preview, [])

    def test_empty_selection_does_not_mean_apply_all(self):
        row = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        preview = self.classifier.preview_rule(self.rule())
        self.assertEqual(self.classifier.apply_rule(preview, []), 0)
        self.assertEqual(self.repo.get_transaction(row["id"]), row)

    def test_changed_or_disabled_rule_invalidates_preview(self):
        row = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        rule = self.rule()
        preview = self.classifier.preview_rule(rule)
        self.rule(category=self.expenses[1])
        with self.assertRaisesRegex(ValueError, "new preview"):
            self.classifier.apply_rule(preview, [row["id"]])
        preview = self.classifier.preview_rule(rule)
        self.classifier.disable_rule(rule)
        with self.assertRaisesRegex(ValueError, "new preview"):
            self.classifier.apply_rule(preview, [row["id"]])
        self.assertEqual(self.repo.get_transaction(row["id"]), row)

    def test_manual_override_after_preview_rolls_back_entire_selected_batch(self):
        rows = self.import_rows(["TEST PURCHASE|Store 12"] * 2)
        preview = self.classifier.preview_rule(self.rule())
        self.classifier.reclassify(rows[1]["id"], self.other)
        before = [self.repo.get_transaction(r["id"]) for r in rows]
        with self.assertRaisesRegex(ValueError, "ineligible"):
            self.classifier.apply_rule(preview, [r["id"] for r in rows])
        self.assertEqual([self.repo.get_transaction(r["id"]) for r in rows], before)

    def test_unselected_change_does_not_block_selected_row(self):
        rows = self.import_rows(["TEST PURCHASE|Store 12"] * 2)
        preview = self.classifier.preview_rule(self.rule())
        self.classifier.reclassify(rows[1]["id"], self.other)
        self.assertEqual(self.classifier.apply_rule(preview, [rows[0]["id"]]), 1)
        self.assertEqual(self.repo.get_transaction(rows[1]["id"])["category_id"], self.other)

    def test_forged_foreign_duplicate_and_consumed_selections_are_rejected(self):
        row = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        preview = self.classifier.preview_rule(self.rule())
        with self.assertRaises(ClassificationBlocked):
            self.classifier.apply_rule(replace(preview), [row["id"]])
        for selection in ([row["id"], row["id"]], [999999]):
            with self.assertRaises(ClassificationBlocked):
                self.classifier.apply_rule(preview, selection)
        self.classifier.apply_rule(preview, [row["id"]])
        with self.assertRaises(ClassificationBlocked):
            self.classifier.apply_rule(preview, [row["id"]])

    def test_effective_direction_change_and_excluded_kind_are_ineligible(self):
        rows = self.import_rows(["TEST PURCHASE|Store 12"] * 2)
        self.repo.edit_transaction(rows[0]["id"], reason="Synthetic direction change", amount_minor=100, manual_lock=False)
        self.repo.edit_transaction(rows[1]["id"], reason="Synthetic exclusion", kind="excluded", manual_lock=False)
        self.assertEqual(self.classifier.preview_rule(self.rule()).candidates, ())

    def test_import_classification_failure_rolls_back_everything(self):
        rule = self.rule()
        preview = self.prepare(["TEST PURCHASE|Store 12"] * 2)
        self.repo._connection.execute("""CREATE TEMP TRIGGER fail_classification BEFORE INSERT ON transaction_edits
            WHEN (SELECT count(*) FROM transaction_edits) > 0
            BEGIN SELECT RAISE(ABORT, 'SYNTHETIC_PRIVATE_MARKER'); END""")
        with self.assertRaises(ImportBlocked) as error:
            self.importer.confirm(preview)
        self.assertNotIn("SYNTHETIC_PRIVATE_MARKER", str(error.exception))
        self.assertEqual(self.counts(), (0, 0, 0, 0))
        self.assertEqual(self.repo.get_rule(rule)["active"], 1)
        self.repo._connection.execute("DROP TRIGGER fail_classification")
        self.assertEqual(self.importer.confirm(preview).imported_count, 2)

    def test_retroactive_audit_failure_rolls_back_all_selected_rows(self):
        rows = self.import_rows(["TEST PURCHASE|Store 12"] * 2)
        preview = self.classifier.preview_rule(self.rule())
        self.repo._connection.execute("""CREATE TEMP TRIGGER fail_retro BEFORE INSERT ON transaction_edits
            WHEN NEW.transaction_id = 2 BEGIN SELECT RAISE(ABORT, 'synthetic failure'); END""")
        with self.assertRaises(ClassificationBlocked):
            self.classifier.apply_rule(preview, [r["id"] for r in rows])
        self.assertEqual([self.repo.get_transaction(r["id"]) for r in rows], rows)

    def test_extractor_failure_is_sanitized_and_rolls_back_import(self):
        preview = self.prepare(["TEST PURCHASE|Store 12"])
        with patch.object(self.classifier, "extractor", side_effect=ValueError("SYNTHETIC_PRIVATE_EXCEPTION")):
            with self.assertRaises(ImportBlocked) as error:
                self.importer.confirm(preview)
        self.assertNotIn("SYNTHETIC_PRIVATE_EXCEPTION", str(error.exception))
        self.assertEqual(self.counts(), (0, 0, 0, 0))

    def test_selected_manual_assignment_rolls_back_incompatible_second_row(self):
        rows = self.import_rows(["Unknown debit", "Unknown credit"], amounts=[-100, 100])
        with self.assertRaises(ValueError):
            self.classifier.assign_category([r["id"] for r in rows], self.expenses[0])
        self.assertEqual([self.repo.get_transaction(r["id"]) for r in rows], rows)

    def test_rule_direction_validation(self):
        for direction, category in (("credit", self.expenses[0]), ("debit", self.income), ("zero", self.expenses[0])):
            with self.subTest(direction=direction), self.assertRaises(ValueError):
                self.rule(direction=direction, category=category)
        self.assertEqual(self.repo.list_rules(), [])

    def test_previews_and_review_repr_do_not_expose_merchant_or_description(self):
        self.import_rows(["TEST PURCHASE|SYNTHETIC_PRIVATE_MERCHANT"])
        queue = self.classifier.review()
        preview = self.classifier.preview_rule(self.rule(merchant="SYNTHETIC_PRIVATE_MERCHANT"))
        rendered = repr(queue) + repr(queue.groups) + repr(preview) + repr(preview.candidates) + repr(preview.selections)
        self.assertNotIn("SYNTHETIC_PRIVATE_MERCHANT", rendered)

    def test_rule_and_applied_id_persist_after_reopening(self):
        rule = self.rule()
        row = self.import_rows(["TEST PURCHASE|Store 12"])[0]
        with Repository(self.path) as reopened:
            self.assertEqual(reopened.get_transaction(row["id"]), row)
            self.assertEqual(reopened.get_rule(rule), self.repo.get_rule(rule))
        with self.assertRaises(sqlite3.IntegrityError):
            self.repo._connection.execute("DELETE FROM merchant_rules WHERE id = ?", (rule,))

    def test_version_two_migration_preserves_data_and_disables_unscoped_rules(self):
        path = self.folder / "version-two.sqlite3"
        with patch.object(db, "SCHEMA_VERSION", 2), patch.object(db, "_CLASSIFICATION_SCHEMA", ()):
            with Repository(path) as old:
                account = old.create_account("000909", "Synthetic legacy")
                old._connection.execute("INSERT INTO merchant_rules(merchant_key,direction,kind,category_id) VALUES (?,?,?,?)",
                                         ("Synthetic Legacy", "debit", "expense", self.expenses[0]))
                source = ImportRow(description="Legacy synthetic", bank_datetime="2026-10-01", occurred_at="2026-10-01",
                                   amount_minor=-100, worksheet="Synthetic", row_number=3)
                old.write_import(ImportInfo(account_id=account, sha256="a" * 64,
                                            parser_version="synthetic", filename="synthetic.xlsx"), [source])
                before = old.query_transactions()[0]
        with Repository(path) as migrated:
            after = migrated.query_transactions()[0]
            self.assertEqual({key: after[key] for key in before}, before)
            self.assertEqual(migrated.list_rules(), [])
            rule = migrated.get_rule(1)
            self.assertEqual((rule["bank"], rule["active"], rule["merchant_key"]), ("", 0, "Synthetic Legacy"))
            self.assertIsNone(after["applied_rule_id"])
            migrated.initialize()
            self.assertEqual(migrated._connection.execute("PRAGMA user_version").fetchone()[0], db.SCHEMA_VERSION)

    def test_failed_classification_migration_rolls_back_schema_and_rules(self):
        path = self.folder / "failed-classification-migration.sqlite3"
        with patch.object(db, "SCHEMA_VERSION", 2), patch.object(db, "_CLASSIFICATION_SCHEMA", ()):
            with Repository(path) as old:
                old._connection.execute("INSERT INTO merchant_rules(merchant_key,direction,kind,category_id) VALUES (?,?,?,?)",
                                         ("Synthetic Legacy", "debit", "expense", self.expenses[0]))
        with patch.object(db, "_CLASSIFICATION_SCHEMA", db._CLASSIFICATION_SCHEMA + ("INVALID SYNTHETIC SQL",)):
            with self.assertRaises(sqlite3.OperationalError):
                Repository(path)
        connection = sqlite3.connect(path)
        try:
            self.assertEqual(connection.execute("PRAGMA user_version").fetchone()[0], 2)
            self.assertEqual(connection.execute("SELECT merchant_key FROM merchant_rules").fetchone()[0], "Synthetic Legacy")
            columns = {row[1] for row in connection.execute("PRAGMA table_info(transactions)")}
            self.assertNotIn("applied_rule_id", columns)
        finally:
            connection.close()


class KhanMerchantExtractionTests(unittest.TestCase):
    """User-confirmed structure, entirely fictional descriptions and accounts."""

    def setUp(self):
        root = Path(__file__).resolve().parents[1]
        temp_root = root / "data" / "test-temp"
        temp_root.mkdir(parents=True, exist_ok=True)
        temporary = tempfile.TemporaryDirectory(prefix="khan-merchant-test-", dir=temp_root)
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        self.repo = Repository(self.folder / "synthetic.sqlite3")
        self.addCleanup(self.repo.close)
        self.account = self.repo.create_account("MN00000000000000000000", "Synthetic Khan")
        self.classifier = ClassificationService(self.repo)
        self.expenses = [c["id"] for c in self.repo.list_categories() if c["direction"] == "expense"]
        self.income = next(c["id"] for c in self.repo.list_categories() if c["direction"] == "income")
        self.sequence = 0

    def extract(self, description):
        return extract_merchant("Khan Bank", "khan-mnt-1", description)

    def import_rows(self, descriptions, amounts=None):
        self.sequence += 1
        amounts = [-100] * len(descriptions) if amounts is None else amounts
        sources = [ImportRow(description=description, bank_datetime=f"2026-10-10 10:00:{i:02d}",
                             occurred_at=f"2026-10-10T10:00:{i:02d}", amount_minor=amount,
                             worksheet="Synthetic", row_number=i + 3)
                   for i, (description, amount) in enumerate(zip(descriptions, amounts))]
        identifier = self.repo.write_import(
            ImportInfo(account_id=self.account, sha256=f"{self.sequence:064x}",
                       parser_version="khan-mnt-1", filename="synthetic.xlsx"),
            sources, classify=self.classifier.classify_import,
        )
        return [self.repo.get_transaction(row["id"]) for row in
                sorted(self.repo.query_transactions(import_id=identifier), key=lambda r: r["id"])]

    def rule(self, text, category=None):
        return self.classifier.save_rule("Khan Bank", text, "debit", category or self.expenses[0])

    def test_confirmed_structures_preserve_complete_remainder(self):
        examples = (
            ("A1B2C3", "EXAMPLE GAME SERVICES    SINGAPORE"),
            ("D4E5F6", "SAMPLE SOFTWARE* MONTHLY SUBS SAMPLE.COM"),
            ("G7H8J9", "DEMO ONLINE STORE 123456789 DEMO.MERCHANT"),
        )
        for code, remainder in examples:
            with self.subTest(code=code):
                match = self.extract(code + " " + remainder)
                self.assertEqual(match.display_text, remainder)
                self.assertEqual(match.key, normalize_merchant_key(remainder))
                self.assertEqual((match.pattern, match.transaction_type), ("khan-card-auth6-v1", "purchase"))

    def test_only_one_authorization_token_is_removed(self):
        remainder = "Z9Y8X7 DEMO STORE 42 DEMO.EXAMPLE  "
        for code in ("A1B2C3", "d4e5f6", "ABCDEF", "123456"):
            match = self.extract(code + "\t " + remainder)
            self.assertEqual(match.display_text, remainder)
            self.assertEqual(match.key, "z9y8x7 demo store 42 demo.example")

    def test_unicode_whitespace_normalization_and_distinct_suffixes(self):
        first = self.extract("A1B2C3 Cafe\u0301\tЖишээ\u00a012    FICTIONLAND  ")
        second = self.extract("D4E5F6 CAFÉ ЖИШЭЭ 12 FICTIONLAND")
        self.assertEqual(first.key, second.key)
        self.assertEqual(first.display_text, "Cafe\u0301\tЖишээ\u00a012    FICTIONLAND  ")
        suffixes = ("DEMO 12", "DEMO 13", "DEMO-12", "DEMO.12", "DEMO １２", "DEMO 12 WEST",
                    "DEMO 12 FICTIONLAND", "DEMO 12 DEMO.EXAMPLE", "DEMO 12 OTHER.EXAMPLE")
        self.assertEqual(len({self.extract("A1B2C3 " + text).key for text in suffixes}), len(suffixes))

    def test_unsupported_scope_and_malformed_structure_abstain(self):
        self.assertIsNone(extract_merchant("Other Bank", "khan-mnt-1", "A1B2C3 DEMO STORE"))
        self.assertIsNone(extract_merchant("Khan Bank", "unknown-version", "A1B2C3 DEMO STORE"))
        for description in ("", "A1B2C3", "A1B2C3   ", "A1B2C DEMO STORE", "A1B2C34 DEMO STORE",
                            "A1_B2C DEMO STORE", "Ａ1B2C3 DEMO STORE", "A1B2C3:DEMO STORE",
                            " A1B2C3 DEMO STORE", "A1B2C3 DEMO\nSTORE", "A1B2C3 DEMO\rSTORE",
                            "A1B2C3 DEMO\x00STORE", "POS DEMO STORE", "ПОС DEMO STORE"):
            with self.subTest(description=description):
                self.assertIsNone(self.extract(description))

    def test_ambiguous_non_purchase_markers_never_apply_rules(self):
        descriptions = ["QPay DEMO STORE", "TRANSFER DEMO STORE", "FEE DEMO STORE", "REFUND DEMO STORE"]
        for marker in ("QPay", "TRANSFER", "FEE", "FEES", "REFUND", "REVERSAL", "CREDIT", "INCOME",
                       "ATM", "Шилжүүлэг", "Шимтгэл", "Буцаалт", "Орлого"):
            remainder = marker + ": DEMO STORE"
            self.rule(remainder)
            descriptions.append("A1B2C3 " + remainder)
        rows = self.import_rows(descriptions)
        self.assertTrue(all(row["merchant_key"] is None and row["category_id"] is None for row in rows))
        self.assertEqual(len(self.classifier.review().individual), len(rows))

    def test_repeated_debits_match_but_credits_and_zero_remain_unclassified(self):
        rule = self.rule("DEMO STORE 12 FICTIONLAND")
        self.classifier.save_rule("Khan Bank", "DEMO STORE 12 FICTIONLAND", "credit", self.income)
        descriptions = [code + " DEMO STORE 12 FICTIONLAND" for code in ("A1B2C3", "D4E5F6", "G7H8J9", "K1L2M3")]
        rows = self.import_rows(descriptions, [-100, -200, 100, 0])
        self.assertEqual([r["applied_rule_id"] for r in rows], [rule, rule, None, None])
        self.assertTrue(all(r["merchant_key"] is None for r in rows[2:]))
        self.assertEqual([r["original_description"] for r in rows], descriptions)
        self.assertEqual([r["description"] for r in rows], descriptions)
        self.assertTrue(all(r["original_reference"] is None for r in rows))

    def test_unknown_merchants_group_by_entire_remainder_without_categories(self):
        rows = self.import_rows(["A1B2C3 DEMO STORE 12", "D4E5F6 demo store 12", "G7H8J9 DEMO STORE 13", "QPay DEMO STORE"])
        queue = self.classifier.review()
        self.assertEqual(sorted(len(group.items) for group in queue.groups), [1, 2])
        self.assertEqual([item.transaction_id for item in queue.individual], [rows[-1]["id"]])
        self.assertTrue(all(row["category_id"] is None for row in rows))

    def test_real_parser_import_and_manual_override_survive_reimport(self):
        source = Path(__file__).parent / "fixtures" / "khan_synthetic.xlsx"
        path = self.folder / "fictional-card.xlsx"
        book = load_workbook(source)
        description = "A1B2C3 DEMO SOFTWARE* MONTHLY SUBS DEMO.EXAMPLE"
        try:
            book.active["G4"] = description
            book.save(path)
        finally:
            book.close()
        rule = self.rule("DEMO SOFTWARE* MONTHLY SUBS DEMO.EXAMPLE")
        importer = ImportService(self.repo)
        preview = importer.preview(path, self.account)
        self.assertTrue(preview.is_valid)
        result = importer.confirm(preview)
        row = next(self.repo.get_transaction(r["id"]) for r in self.repo.query_transactions(import_id=result.import_id)
                   if r["description"] == description)
        self.assertEqual(row["applied_rule_id"], rule)
        self.assertEqual(row["original_description"], description)
        self.classifier.reclassify(row["id"], self.expenses[1])
        self.rule("DEMO SOFTWARE* MONTHLY SUBS DEMO.EXAMPLE", self.expenses[2])
        self.assertEqual(importer.confirm(importer.preview(path, self.account)).status, "identical_file")
        # Changed workbook bytes with the same source transactions require explicit overlap decisions.
        book = load_workbook(path)
        try:
            book.active.title = "Different synthetic provenance"
            book.save(path)
        finally:
            book.close()
        overlap = importer.preview(path, self.account)
        self.assertTrue(overlap.groups)
        importer.confirm(overlap, {r.index: "skip" for r in overlap.rows if r.status == "ambiguous"})
        current = self.repo.get_transaction(row["id"])
        self.assertEqual((current["category_id"], current["manual_lock"]), (self.expenses[1], 1))
        self.assertEqual(current["original_description"], description)
        self.assertEqual(current["dedup_key"], row["dedup_key"])
        self.classifier.disable_rule(rule)
        future = self.import_rows(["D4E5F6 DEMO SOFTWARE* MONTHLY SUBS DEMO.EXAMPLE"])[0]
        self.assertIsNone(future["category_id"])

    def test_retroactive_application_selects_only_unlocked_eligible_purchases(self):
        rows = self.import_rows(["A1B2C3 DEMO STORE", "D4E5F6 DEMO STORE", "G7H8J9 DEMO STORE",
                                 "K1L2M3 DEMO STORE", "N4P5Q6 DEMO STORE", "R7S8T9 DEMO STORE"],
                                [-100, -100, -100, -100, 100, -100])
        self.classifier.reclassify(rows[2]["id"], self.expenses[1])
        self.repo.edit_transaction(rows[3]["id"], reason="Synthetic transfer", kind="transfer", manual_lock=False)
        self.repo.edit_transaction(rows[5]["id"], reason="Synthetic exclusion", kind="excluded", manual_lock=False)
        rule = self.rule("DEMO STORE")
        preview = self.classifier.preview_rule(rule)
        self.assertEqual({r.transaction_id for r in preview.candidates}, {rows[0]["id"], rows[1]["id"]})
        self.assertEqual(self.classifier.apply_rule(preview, [rows[0]["id"]]), 1)
        self.assertEqual(self.repo.get_transaction(rows[0]["id"])["applied_rule_id"], rule)
        self.assertIsNone(self.repo.get_transaction(rows[1]["id"])["category_id"])
        self.assertEqual(self.repo.get_transaction(rows[2]["id"])["category_id"], self.expenses[1])


if __name__ == "__main__":
    unittest.main()
