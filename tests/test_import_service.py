"""Import acceptance: synthetic workbooks/records and workspace databases only."""

import hashlib
import sqlite3
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app import db
from app.db import Repository
from app.import_service import ImportBlocked, ImportService, VerifiedReferencePolicy
from app.models import ImportInfo, ImportRow
from app.parsers.base import ParseIssue, ParseResult, ParsedTransaction, StatementMetadata
from app.parsers.khan import KhanParser


class SyntheticParser:
    """A test-only bank contract, not evidence of Khan reference semantics."""

    version = "synthetic-unique-references-1"

    def parse_bytes(self, payload: bytes) -> ParseResult:
        self.seen_payload = payload
        return self.result


class ImportServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        root = Path(__file__).resolve().parents[1]
        temp_root = root / "data" / "test-temp"
        temp_root.mkdir(parents=True, exist_ok=True)
        temporary = tempfile.TemporaryDirectory(prefix="import-test-", dir=temp_root)
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        self.path = self.folder / "synthetic.sqlite3"
        self.repo = Repository(self.path)
        self.addCleanup(self.repo.close)
        self.account = self.repo.create_account("000001", "Synthetic account", bank="Synthetic Test Bank")
        self.khan_account = self.repo.create_account("MN00000000000000000000", "Synthetic Khan account")
        self.parser = SyntheticParser()
        self.service = ImportService(self.repo, self.parser)
        self.sequence = 0
        self.fixture = root / "tests" / "fixtures" / "khan_synthetic.xlsx"

    def row(self, **changes) -> ImportRow:
        return ImportRow(**{
            "description": "Synthetic repeated row", "bank_datetime": "2026-10-05 09:10:11",
            "occurred_at": "2026-10-05T09:10:11", "amount_minor": -250,
            "worksheet": "Synthetic", "row_number": 3, **changes,
        })

    def prepare(self, rows, *, account=None, issues=(), candidate_count=None):
        account = self.account if account is None else account
        selected = next(a for a in self.repo.list_accounts() if a["id"] == account)
        self.parser.result = ParseResult(
            metadata=StatementMetadata(parser_version=self.parser.version, currency=selected["currency"],
                                       account_identifier=selected["identifier"], worksheet="Synthetic",
                                       period_start="2026-10-01", period_end="2026-10-31", printed_on="2026-11-01"),
            records=tuple(ParsedTransaction(source=r, branch="0001", counterparty_account=None,
                                            opening_balance_minor=0, credit_minor=max(r.amount_minor, 0),
                                            debit_minor=min(r.amount_minor, 0)) for r in rows),
            issues=issues, candidate_row_count=len(rows) if candidate_count is None else candidate_count,
        )
        self.sequence += 1
        path = self.folder / f"synthetic-{self.sequence}.xlsx"
        path.write_bytes(f"INVENTED TEST PAYLOAD {self.sequence}".encode())
        return path

    def preview(self, rows, **kwargs):
        path = self.prepare(rows, **kwargs)
        return self.service.preview(path, kwargs.get("account", self.account))

    def enable_verified_ids(self):
        self.service = ImportService(self.repo, self.parser, reference_policy=VerifiedReferencePolicy(
            bank="Synthetic Test Bank", parser_version=self.parser.version,
            namespace="synthetic-test-only", evidence="Test contract: references are unique per account and stable across exports.",
        ))

    def counts(self):
        return tuple(self.repo._connection.execute(f"SELECT count(*) FROM {name}").fetchone()[0]
                     for name in ("imports", "transactions", "import_decisions"))

    def test_first_import_real_parser_on_synthetic_workbook(self):
        service = ImportService(self.repo)
        before = self.path.read_bytes()
        preview = service.preview(self.fixture, self.khan_account)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertTrue(preview.is_valid)
        self.assertEqual((preview.valid_count, preview.invalid_count, preview.duplicate_count, preview.ambiguous_count), (3, 0, 0, 0))
        result = service.confirm(preview)
        self.assertEqual((result.status, result.imported_count, result.skipped_count), ("committed", 3, 0))
        metadata = self.repo.get_import(result.import_id)
        self.assertEqual(metadata["sha256"], hashlib.sha256(self.fixture.read_bytes()).hexdigest())
        self.assertEqual(metadata["parser_version"], "khan-mnt-1")
        self.assertEqual(metadata["total_count"], 3)
        transactions = self.repo.query_transactions(import_id=result.import_id)
        self.assertTrue(all(t["kind"] == "unclassified" and t["category_id"] is None for t in transactions))
        self.assertEqual(len(self.repo.get_import_decisions(result.import_id)), 3)

    def test_identical_file_twice_is_no_write_even_renamed(self):
        service = ImportService(self.repo)
        first = service.confirm(service.preview(self.fixture, self.khan_account))
        renamed = self.folder / "renamed.xlsx"
        renamed.write_bytes(self.fixture.read_bytes())
        before = self.path.read_bytes()
        preview = service.preview(renamed, self.khan_account)
        self.assertEqual((preview.valid_count, preview.duplicate_count), (0, 3))
        second = service.confirm(preview)
        self.assertEqual((second.status, second.import_id, second.imported_count), ("identical_file", first.import_id, 0))
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(self.counts(), (1, 3, 3))

    def test_hash_dedup_is_scoped_to_account(self):
        path = self.prepare([self.row()])
        self.service.confirm(self.service.preview(path, self.account))
        other = self.repo.create_account("000002", "Synthetic second", bank="Synthetic Test Bank")
        self.parser.result = replace(self.parser.result, metadata=replace(self.parser.result.metadata, account_identifier="000002"))
        preview = self.service.preview(path, other)
        self.assertIsNone(preview.identical_import_id)
        self.assertEqual(self.service.confirm(preview).imported_count, 1)

    def test_khan_overlap_changed_filename_sheet_and_rows_requires_decisions(self):
        service = ImportService(self.repo)
        service.confirm(service.preview(self.fixture, self.khan_account))
        path = self.folder / "synthetic-overlap.xlsx"
        book = load_workbook(self.fixture)
        try:
            book.active.title = "Renamed synthetic sheet"
            book.active["B1"] = "10/21/2026"
            book.active.insert_rows(1)
            book.save(path)
        finally:
            book.close()
        preview = service.preview(path, self.khan_account)
        self.assertTrue(preview.is_valid)
        self.assertEqual((preview.valid_count, preview.duplicate_count, preview.ambiguous_count), (0, 0, 3))
        result = service.confirm(preview, {0: "skip", 1: "skip", 2: "skip"})
        self.assertEqual((result.imported_count, result.skipped_count), (0, 3))
        self.assertEqual(self.counts(), (2, 3, 6))

    def test_overlapping_statements_with_test_verified_ids(self):
        self.enable_verified_ids()
        rows = [self.row(reference="TEST-A"), self.row(reference="TEST-B", row_number=4)]
        self.service.confirm(self.preview(rows))
        preview = self.preview([replace(rows[1], worksheet="Changed sheet", row_number=9),
                                self.row(reference="TEST-C", row_number=10)])
        self.assertEqual((preview.valid_count, preview.duplicate_count, preview.ambiguous_count), (1, 1, 0))
        result = self.service.confirm(preview)
        self.assertEqual((result.imported_count, result.skipped_count), (1, 1))
        decisions = self.repo.get_import_decisions(result.import_id)
        self.assertEqual((decisions[0]["action"], decisions[0]["reason"]), ("skip", "verified_reference"))
        self.assertEqual(self.counts(), (2, 3, 4))

    def test_first_import_preserves_legitimate_identical_occurrences(self):
        preview = self.preview([self.row(), self.row(row_number=4)])
        self.assertEqual(preview.valid_count, 2)
        self.assertEqual(self.service.confirm(preview).imported_count, 2)
        rows = self.repo.query_transactions(account_id=self.account)
        self.assertEqual(rows[0]["dedup_key"], rows[1]["dedup_key"])
        self.assertNotEqual(rows[0]["source_id"], rows[1]["source_id"])

    def test_ambiguous_group_requires_complete_explicit_decisions_and_persists_them(self):
        self.service.confirm(self.preview([self.row()]))
        preview = self.preview([self.row(row_number=8), self.row(row_number=9)])
        self.assertEqual((preview.valid_count, preview.ambiguous_count), (0, 2))
        self.assertEqual(preview.groups[0].row_indexes, (0, 1))
        self.assertEqual(len(preview.groups[0].existing), 1)
        before = self.counts()
        for decisions in (None, {}, {0: "skip"}, {0: "skip", 1: "skip"}, {0: "skip", 1: "keep", 2: "keep"}):
            with self.subTest(decisions=decisions), self.assertRaises(ImportBlocked):
                self.service.confirm(preview, decisions)
            self.assertEqual(self.counts(), before)
        result = self.service.confirm(preview, {0: "skip", 1: "keep"})
        self.assertEqual((result.imported_count, result.skipped_count), (1, 1))
        with Repository(self.path) as reopened:
            audit = reopened.get_import_decisions(result.import_id)
            self.assertEqual([d["action"] for d in audit], ["skip", "keep"])
            self.assertTrue(all(d["reason"] == "manual" for d in audit))
            self.assertEqual([d["source"]["row_number"] for d in audit], [8, 9])

    def test_keep_all_ambiguous_occurrences_is_allowed(self):
        self.service.confirm(self.preview([self.row()]))
        preview = self.preview([self.row(), self.row(row_number=4)])
        result = self.service.confirm(preview, {0: "keep", 1: "keep"})
        self.assertEqual(result.imported_count, 2)
        self.assertEqual(self.counts()[1], 3)

    def test_skip_all_records_still_persists_import_and_decisions(self):
        self.service.confirm(self.preview([self.row()]))
        preview = self.preview([self.row()])
        result = self.service.confirm(preview, {0: "skip"})
        self.assertEqual((result.imported_count, result.skipped_count), (0, 1))
        self.assertEqual(self.counts(), (2, 1, 2))
        self.assertEqual(self.repo.get_import(result.import_id)["total_count"], 1)

    def test_malformed_workbook_is_blocked_without_partial_import(self):
        path = self.folder / "malformed.xlsx"
        path.write_bytes(b"synthetic invalid zip")
        service = ImportService(self.repo)
        preview = service.preview(path, self.khan_account)
        self.assertFalse(preview.is_valid)
        with self.assertRaises(ImportBlocked):
            service.confirm(preview)
        self.assertEqual(self.counts(), (0, 0, 0))

    def test_invalid_row_blocks_good_rows_with_real_parser(self):
        path = self.folder / "bad-row.xlsx"
        book = load_workbook(self.fixture)
        try:
            book.active["D3"] = "invalid synthetic amount"
            book.save(path)
        finally:
            book.close()
        service = ImportService(self.repo)
        preview = service.preview(path, self.khan_account)
        self.assertGreater(preview.valid_count, 0)
        self.assertGreater(preview.invalid_count, 0)
        with self.assertRaises(ImportBlocked):
            service.confirm(preview)
        self.assertEqual(self.counts(), (0, 0, 0))

    def test_global_reconciliation_error_blocks_even_without_invalid_rows(self):
        preview = self.preview([self.row()], issues=(ParseIssue("total_mismatch", "Synthetic reconciliation failed."),))
        self.assertEqual(preview.invalid_count, 0)
        self.assertFalse(preview.is_valid)
        with self.assertRaises(ImportBlocked):
            self.service.confirm(preview)
        self.assertEqual(self.counts(), (0, 0, 0))

    def test_canceled_preview_never_writes_database_or_source(self):
        path = self.prepare([self.row()])
        database_before, source_before = self.path.read_bytes(), path.read_bytes()
        preview = self.service.preview(path, self.account)
        self.service.cancel(preview)
        self.assertEqual(self.path.read_bytes(), database_before)
        self.assertEqual(path.read_bytes(), source_before)
        with self.assertRaises(ImportBlocked):
            self.service.confirm(preview)

    def test_injected_row_failure_rolls_back_metadata_and_previous_rows(self):
        preview = self.preview([self.row(), self.row(row_number=4)])
        self.repo._connection.execute("""CREATE TEMP TRIGGER fail_row BEFORE INSERT ON transactions
            WHEN NEW.row_number = 4 BEGIN SELECT RAISE(ABORT, 'SYNTHETIC_PRIVATE_VALUE'); END""")
        with self.assertRaises(ImportBlocked) as error:
            self.service.confirm(preview)
        self.assertNotIn("SYNTHETIC_PRIVATE_VALUE", str(error.exception))
        self.assertEqual(self.counts(), (0, 0, 0))
        self.repo._connection.execute("DROP TRIGGER fail_row")
        self.assertEqual(self.service.confirm(preview).imported_count, 2)

    def test_injected_decision_failure_rolls_back_rows_and_metadata(self):
        preview = self.preview([self.row(), self.row(row_number=4)])
        self.repo._connection.execute("""CREATE TEMP TRIGGER fail_decision BEFORE INSERT ON import_decisions
            WHEN NEW.row_index = 1 BEGIN SELECT RAISE(ABORT, 'synthetic failure'); END""")
        with self.assertRaises(ImportBlocked):
            self.service.confirm(preview)
        self.assertEqual(self.counts(), (0, 0, 0))
        with Repository(self.path) as reopened:
            self.assertEqual(reopened.query_transactions(), [])
            self.assertEqual(reopened.get_import_state(self.account)["imports"], [])

    def test_manually_edited_transaction_matches_original_without_overwrite(self):
        self.service.confirm(self.preview([self.row()]))
        transaction = self.repo.query_transactions(account_id=self.account)[0]
        category = next(c["id"] for c in self.repo.list_categories() if c["direction"] == "expense")
        self.repo.assign_category(transaction["id"], category, reason="Synthetic manual category")
        self.repo.edit_transaction(transaction["id"], reason="Synthetic correction", description="Edited synthetic text",
                                   amount_minor=-999, occurred_at="2026-10-09T01:02:03", reference="EDITED", balance_minor=888,
                                   merchant_key="synthetic merchant", note="Synthetic note")
        before = self.repo.get_transaction(transaction["id"])
        history = self.repo.get_history(transaction["id"])
        preview = self.preview([self.row(worksheet="New provenance", row_number=20)])
        self.assertEqual(preview.ambiguous_count, 1)
        self.assertEqual(preview.rows[0].candidates[0].source.description, "Synthetic repeated row")
        self.service.confirm(preview, {0: "skip"})
        self.assertEqual(self.repo.get_transaction(transaction["id"]), before)
        self.assertEqual(self.repo.get_history(transaction["id"]), history)

    def test_edits_after_preview_do_not_stale_or_overwrite_category(self):
        self.service.confirm(self.preview([self.row()]))
        preview = self.preview([self.row()])
        transaction = self.repo.query_transactions(account_id=self.account)[0]
        self.repo.edit_transaction(transaction["id"], reason="Synthetic exclusion", kind="excluded", note="Manual")
        self.service.confirm(preview, {0: "skip"})
        self.assertEqual(self.repo.get_transaction(transaction["id"])["kind"], "excluded")

    def test_concurrent_import_stales_preview_and_rolls_back(self):
        preview = self.preview([self.row()])
        with Repository(self.path) as other:
            other.write_import(ImportInfo(account_id=self.account, sha256="a" * 64,
                                          parser_version="synthetic", filename="synthetic.xlsx"), [self.row()])
        with self.assertRaisesRegex(ImportBlocked, "new preview"):
            self.service.confirm(preview)
        self.assertEqual(self.counts(), (1, 1, 0))

    def test_revision_is_checked_inside_the_atomic_write(self):
        preview = self.preview([self.row()])
        original_write = self.repo.write_import
        def concurrent_write(*args, **kwargs):
            with Repository(self.path) as other:
                other.write_import(ImportInfo(account_id=self.account, sha256="b" * 64,
                                              parser_version="synthetic", filename="synthetic.xlsx"), [self.row()])
            return original_write(*args, **kwargs)
        with patch.object(self.repo, "write_import", side_effect=concurrent_write):
            with self.assertRaisesRegex(ImportBlocked, "new preview"):
                self.service.confirm(preview)
        self.assertEqual(self.counts(), (1, 1, 0))

    def test_preview_cannot_be_forged_or_confirmed_twice(self):
        preview = self.preview([self.row()])
        with self.assertRaises(ImportBlocked):
            self.service.confirm(replace(preview, invalid_count=0))
        with self.assertRaises(ImportBlocked):
            ImportService(self.repo, self.parser).confirm(preview)
        self.service.confirm(preview)
        with self.assertRaises(ImportBlocked):
            self.service.confirm(preview)
        self.assertEqual(self.counts(), (1, 1, 1))

    def test_hash_and_parse_use_same_bytes_without_temp_copy(self):
        path = self.prepare([self.row()])
        original_bytes = path.read_bytes()
        parse_bytes = self.parser.parse_bytes
        def change_source(payload):
            path.write_bytes(b"DIFFERENT INVENTED CONTENT")
            return parse_bytes(payload)
        with patch.object(self.parser, "parse_bytes", side_effect=change_source):
            preview = self.service.preview(path, self.account)
        self.assertEqual(self.parser.seen_payload, original_bytes)
        result = self.service.confirm(preview)
        self.assertEqual(self.repo.get_import(result.import_id)["sha256"], hashlib.sha256(original_bytes).hexdigest())

    def test_wrong_selected_account_is_blocked(self):
        service = ImportService(self.repo)
        preview = service.preview(self.fixture, self.account)
        self.assertIn("account_mismatch", {i.code for i in preview.issues})
        with self.assertRaises(ImportBlocked):
            service.confirm(preview)

    def test_unverified_reference_never_automatically_skips(self):
        self.service.confirm(self.preview([self.row(reference="TEST-A")]))
        preview = self.preview([self.row(reference="TEST-A")])
        self.assertEqual((preview.duplicate_count, preview.ambiguous_count), (0, 1))

    def test_reference_collision_under_verified_contract_is_invalid(self):
        self.enable_verified_ids()
        preview = self.preview([self.row(reference="TEST-A"), self.row(reference="TEST-A", row_number=4)])
        self.assertEqual(preview.invalid_count, 2)
        with self.assertRaises(ImportBlocked):
            self.service.confirm(preview)
        self.assertEqual(self.counts(), (0, 0, 0))

    def test_verified_reference_with_changed_source_requires_decision(self):
        self.enable_verified_ids()
        self.service.confirm(self.preview([self.row(reference="TEST-A")]))
        preview = self.preview([self.row(reference="TEST-A", amount_minor=-999)])
        self.assertEqual(preview.ambiguous_count, 1)
        self.service.confirm(preview, {0: "keep"})
        self.assertEqual(self.counts()[1], 2)

    def test_verified_reference_is_not_retroactively_assumed_for_legacy_import(self):
        self.service.confirm(self.preview([self.row(reference="TEST-A")]))
        self.enable_verified_ids()
        preview = self.preview([self.row(reference="TEST-A")])
        self.assertEqual((preview.ambiguous_count, preview.duplicate_count), (1, 0))

    def test_optional_values_and_datetime_precision_affect_identity(self):
        self.service.confirm(self.preview([self.row(balance_minor=1000, reference="TEST-A")]))
        cases = [
            ({"balance_minor": 2000, "reference": "TEST-A"}, "valid"),
            ({"balance_minor": 1000, "reference": "TEST-B"}, "valid"),
            ({"balance_minor": None, "reference": "TEST-A"}, "ambiguous"),
            ({"balance_minor": 1000, "reference": None}, "ambiguous"),
            ({"balance_minor": 1000, "reference": "TEST-A", "occurred_at": "2026-10-05"}, "valid"),
            ({"balance_minor": 1000, "reference": "TEST-A", "description": " Synthetic repeated row"}, "valid"),
        ]
        for fields, expected in cases:
            with self.subTest(fields=fields):
                preview = self.preview([self.row(**fields)])
                self.assertEqual(preview.rows[0].status, expected)
                self.service.cancel(preview)

    def test_skip_matching_preserves_multiplicity_with_partial_optional_fields(self):
        self.service.confirm(self.preview([self.row(balance_minor=100), self.row(balance_minor=200, row_number=4)]))
        preview = self.preview([self.row(balance_minor=None), self.row(balance_minor=100, row_number=4)])
        result = self.service.confirm(preview, {0: "skip", 1: "skip"})
        self.assertEqual((result.imported_count, result.skipped_count), (0, 2))

    def test_matching_uses_all_rows_beyond_query_page_limit(self):
        rows = [self.row(description=f"Synthetic item {i}", row_number=i + 3) for i in range(1005)]
        self.service.confirm(self.preview(rows))
        preview = self.preview([rows[0], rows[-1]])
        self.assertEqual(preview.ambiguous_count, 2)

    def test_preview_representations_hide_sensitive_source_fields(self):
        row = self.row(description="SYNTHETIC_PRIVATE_DESCRIPTION", reference="SYNTHETIC_PRIVATE_REFERENCE")
        self.service.confirm(self.preview([row]))
        preview = self.preview([row])
        text = repr(preview) + repr(preview.rows) + repr(preview.groups)
        for marker in (row.description, row.reference, row.worksheet, row.bank_datetime):
            self.assertNotIn(marker, text)

    def test_decision_history_is_immutable(self):
        self.service.confirm(self.preview([self.row()]))
        with self.assertRaises(sqlite3.IntegrityError):
            self.repo._connection.execute("UPDATE import_decisions SET action = 'skip'")
        with self.assertRaises(sqlite3.IntegrityError):
            self.repo._connection.execute("DELETE FROM import_decisions")

    def test_schema_v1_migration_preserves_existing_data_and_is_idempotent(self):
        old_path = self.folder / "version-one.sqlite3"
        with patch.object(db, "SCHEMA_VERSION", 1), patch.object(db, "_IMPORT_SCHEMA", ()):
            with Repository(old_path) as old:
                account = old.create_account("000009", "Synthetic legacy")
                old.write_import(ImportInfo(account_id=account, sha256="c" * 64,
                                            parser_version="synthetic", filename="legacy.xlsx"), [self.row()])
                before = old.query_transactions()
                categories = old.list_categories()
        with Repository(old_path) as migrated:
            self.assertEqual(migrated.query_transactions(), before)
            self.assertEqual(migrated.list_categories(), categories)
            self.assertEqual(migrated.get_import_decisions(1), [])
            migrated.initialize()
            self.assertEqual(migrated._connection.execute("PRAGMA user_version").fetchone()[0], 2)

    def test_failed_schema_migration_is_atomic(self):
        old_path = self.folder / "failed-migration.sqlite3"
        with patch.object(db, "SCHEMA_VERSION", 1), patch.object(db, "_IMPORT_SCHEMA", ()):
            with Repository(old_path):
                pass
        with patch.object(db, "_IMPORT_SCHEMA", db._IMPORT_SCHEMA + ("INVALID SYNTHETIC SQL",)):
            with self.assertRaises(sqlite3.OperationalError):
                Repository(old_path)
        connection = sqlite3.connect(old_path)
        try:
            self.assertEqual(connection.execute("PRAGMA user_version").fetchone()[0], 1)
            self.assertIsNone(connection.execute("SELECT name FROM sqlite_master WHERE name = 'import_decisions'").fetchone())
        finally:
            connection.close()


if __name__ == "__main__":
    unittest.main()
